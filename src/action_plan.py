"""
Nubra Trade Replay — Action Plan Document Generation Service.
Generates professionally formatted, evidence-backed, personalized PDF and DOCX
30-Day Trading Behaviour Action Plans.
"""

import io
import os
import html
from dataclasses import dataclass
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple, Union

import pandas as pd
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from src.models import Finding, SummaryMetrics
from src.rules import RULES_CATALOG, get_rule_by_id


# -----------------------------------------------------------------------------
# FONT REGISTRATION & REPORTLAB NUMBERED CANVAS
# -----------------------------------------------------------------------------

_FONT_SETUP_DONE = False
_PDF_FONT_NAME = "Helvetica"
_PDF_FONT_BOLD = "Helvetica-Bold"
_HAS_RUPEE_FONT = False


def init_pdf_fonts() -> Tuple[str, str, bool]:
    """Register unicode-compatible fonts for ReportLab (supporting Rupee symbol ₹)."""
    global _FONT_SETUP_DONE, _PDF_FONT_NAME, _PDF_FONT_BOLD, _HAS_RUPEE_FONT
    if _FONT_SETUP_DONE:
        return _PDF_FONT_NAME, _PDF_FONT_BOLD, _HAS_RUPEE_FONT

    font_candidates = [
        # macOS Helvetica TTC (contains native 0x20B9 Rupee glyph)
        ("/System/Library/Fonts/Helvetica.ttc", "/System/Library/Fonts/Helvetica.ttc", 0, 1),
        # macOS Geneva (contains 0x20B9)
        ("/System/Library/Fonts/Geneva.ttf", None, None, None),
        # Linux DejaVu Sans
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", None, None),
        ("/usr/share/fonts/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf", None, None),
        # macOS Arial Unicode / standard Arial fallbacks
        ("/Library/Fonts/Arial Unicode.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf", None, None),
        ("/System/Library/Fonts/Supplemental/Arial Unicode.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf", None, None),
        ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf", None, None),
    ]

    for norm_path, bold_path, norm_sub, bold_sub in font_candidates:
        if os.path.exists(norm_path):
            try:
                kwargs_norm = {"subfontIndex": norm_sub} if norm_sub is not None else {}
                font_norm = TTFont("NubraFont", norm_path, **kwargs_norm)
                has_rupee = (0x20B9 in font_norm.face.charToGlyph)
                pdfmetrics.registerFont(font_norm)

                if bold_path and os.path.exists(bold_path):
                    kwargs_bold = {"subfontIndex": bold_sub} if bold_sub is not None else {}
                    font_bold = TTFont("NubraFont-Bold", bold_path, **kwargs_bold)
                    pdfmetrics.registerFont(font_bold)
                    pdfmetrics.registerFontFamily(
                        "NubraFont",
                        normal="NubraFont",
                        bold="NubraFont-Bold",
                        italic="NubraFont",
                        boldItalic="NubraFont-Bold",
                    )
                    _PDF_FONT_BOLD = "NubraFont-Bold"
                else:
                    pdfmetrics.registerFontFamily(
                        "NubraFont",
                        normal="NubraFont",
                        bold="NubraFont",
                        italic="NubraFont",
                        boldItalic="NubraFont",
                    )
                    _PDF_FONT_BOLD = "NubraFont"

                _PDF_FONT_NAME = "NubraFont"
                _HAS_RUPEE_FONT = has_rupee
                if has_rupee:
                    break
            except Exception:
                continue

    _FONT_SETUP_DONE = True
    return _PDF_FONT_NAME, _PDF_FONT_BOLD, _HAS_RUPEE_FONT


class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas to compute total page count and draw running header/footer."""

    def __init__(self, *args: Any, **kwargs: Any):
        super().__init__(*args, **kwargs)
        self._saved_page_states: List[Dict[str, Any]] = []

    def showPage(self) -> None:
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self) -> None:
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count: int) -> None:
        font_name, _, _ = init_pdf_fonts()
        self.saveState()
        self.setFont(font_name, 8)
        self.setFillColor(colors.HexColor("#64748B"))

        # Running header on pages > 1
        if self._pageNumber > 1:
            self.drawString(54, 800, "Nubra Trade Replay — 30-Day Trading Behaviour Action Plan")
            self.setStrokeColor(colors.HexColor("#E2E8F0"))
            self.setLineWidth(0.75)
            self.line(54, 792, 541, 792)

        # Running footer on all pages
        self.setStrokeColor(colors.HexColor("#E2E8F0"))
        self.setLineWidth(0.75)
        self.line(54, 45, 541, 45)
        self.drawString(54, 32, "Confidential — Trader Self-Monitoring Action Plan")
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(541, 32, page_text)
        self.restoreState()


# -----------------------------------------------------------------------------
# CURRENCY & TEXT FORMATTING HELPERS
# -----------------------------------------------------------------------------

def format_currency_text(val: float, has_rupee_font: bool = True) -> str:
    """Format currency values cleanly with rupee symbol or Rs. fallback."""
    prefix = "₹" if has_rupee_font else "Rs. "
    if abs(val) < 0.001:
        return f"{prefix}0.00"
    sign = "+" if val > 0 else "-"
    return f"{sign}{prefix}{abs(val):,.2f}"


def clean_currency_text(text: Any) -> str:
    """Format negative Indian rupee amounts consistently as -₹2,400.00, never ₹-2,400.00."""
    s = str(text)
    return s.replace("₹-", "-₹").replace("Rs. -", "-Rs. ")


def safe_text(text: str) -> str:
    """Escape XML characters for ReportLab paragraphs and remove accidental duplicate punctuation."""
    escaped = html.escape(str(text), quote=False)
    return escaped.replace("..", ".")


def normalize_rule_for_plan(active_rule: Dict[str, Any]) -> Dict[str, Any]:
    """
    Normalize rule attributes for document generation ensuring consistent operational triggers.
    For Persona B (activity_ceiling): enforces consistent trigger:
    'When the fourth completed trade of the day is reached, pause and reassess before placing another trade.'
    across rule card, checklist, weekly milestones, success-tracking definitions, and any narrative.
    """
    if not active_rule:
        return active_rule
    rule = dict(active_rule)
    if rule.get("id") == "activity_ceiling":
        consistent_trigger = "When the fourth completed trade of the day is reached, pause and reassess before placing another trade."
        rule["trigger"] = consistent_trigger
        rule["summary"] = consistent_trigger
        rule["action"] = "Pause execution and reassess market conditions and strategy fit before placing another trade."
        rule["checklist"] = [
            consistent_trigger,
            "Reassess market conditions and strategy fit to avoid unnecessary additional trades.",
            "Do not increase trade frequency to compensate for slow market movement or drawdowns.",
        ]
    return rule


# -----------------------------------------------------------------------------
# FINDING & EVIDENCE SELECTION HELPERS
# -----------------------------------------------------------------------------

def analyze_review_and_rule(
    findings: List[Finding], active_rule: Dict[str, Any]
) -> Tuple[Optional[Finding], Optional[Finding], str]:
    """
    Determine the relationship between the committed rule and detected findings:
      - matching_finding: finding that suggested the committed rule (if detected)
      - primary_behavioral_finding: top behavioral finding detected in review
      - explanation_mode: 'matched', 'user_custom', or 'clean'
    """
    matching_finding = None
    primary_behavioral = None

    for f in findings:
        if f.detector_id not in ("clean_baseline", "baseline_clean"):
            if primary_behavioral is None:
                primary_behavioral = f
            if f.suggested_rule_id == active_rule["id"]:
                matching_finding = f

    if matching_finding is not None:
        explanation_mode = "matched"
    elif primary_behavioral is not None:
        explanation_mode = "user_custom"
    else:
        explanation_mode = "clean"

    return matching_finding, primary_behavioral, explanation_mode


# -----------------------------------------------------------------------------
# UNIFIED REPORT PIPELINE CONTEXT
# -----------------------------------------------------------------------------

@dataclass
class ReportContext:
    """
    Normalized report context containing all information required to generate
    an action plan document (PDF or DOCX), shared across all data sources
    (seeded sample personas, custom uploaded tradebooks, and custom rule commitments).
    """
    dataset_label: str
    is_synthetic: bool
    metrics: SummaryMetrics
    findings: List[Finding]
    active_rule: Dict[str, Any]
    df_trades: Optional[pd.DataFrame] = None
    generation_date: Optional[datetime] = None
    matching_finding: Optional[Finding] = None
    primary_behavioral: Optional[Finding] = None
    explanation_mode: str = "matched"
    is_small_sample: bool = False
    small_sample_warning: Optional[str] = None


def build_report_context(
    dataset_label: str,
    is_synthetic: bool,
    metrics: SummaryMetrics,
    findings: List[Finding],
    active_rule: Dict[str, Any],
    df_trades: Optional[pd.DataFrame] = None,
    generation_date: Optional[datetime] = None,
) -> ReportContext:
    """
    Constructs a normalized ReportContext object from raw review state.
    Standardizes rule parameters (e.g. 4th-trade trigger for high activity)
    and evaluates the finding-rule relationship.
    """
    normalized_rule = normalize_rule_for_plan(active_rule)
    matching_finding, primary_behavioral, explanation_mode = analyze_review_and_rule(
        findings, normalized_rule
    )
    gen_dt = generation_date or datetime.now()
    is_small = metrics.is_small_sample if metrics else False
    small_warning = (
        f"The reviewed dataset contains only {metrics.total_trades} completed trades (< 10 trades). "
        "Statistical identification of recurring habits requires larger trade samples."
        if is_small else None
    )

    return ReportContext(
        dataset_label=dataset_label,
        is_synthetic=is_synthetic,
        metrics=metrics,
        findings=findings,
        active_rule=normalized_rule,
        df_trades=df_trades,
        generation_date=gen_dt,
        matching_finding=matching_finding,
        primary_behavioral=primary_behavioral,
        explanation_mode=explanation_mode,
        is_small_sample=is_small,
        small_sample_warning=small_warning,
    )


# -----------------------------------------------------------------------------
# PRIMARY EXPORT: CANONICAL PDF GENERATION
# -----------------------------------------------------------------------------

def generate_action_plan_pdf(
    dataset_label: Optional[Union[ReportContext, str]] = None,
    is_synthetic: Optional[bool] = None,
    metrics: Optional[SummaryMetrics] = None,
    findings: Optional[List[Finding]] = None,
    active_rule: Optional[Dict[str, Any]] = None,
    generation_date: Optional[datetime] = None,
    df_trades: Optional[pd.DataFrame] = None,
    context: Optional[ReportContext] = None,
) -> bytes:
    """
    Canonical entry point for PDF 30-Day Action Plan generation.
    Accepts a normalized ReportContext or individual parameters.
    """
    if isinstance(dataset_label, ReportContext):
        ctx = dataset_label
    elif context is not None:
        ctx = context
    else:
        ctx = build_report_context(
            dataset_label=str(dataset_label or ""),
            is_synthetic=bool(is_synthetic),
            metrics=metrics,
            findings=findings or [],
            active_rule=active_rule or {},
            df_trades=df_trades,
            generation_date=generation_date,
        )

    dataset_label = ctx.dataset_label
    is_synthetic = ctx.is_synthetic
    metrics = ctx.metrics
    findings = ctx.findings
    active_rule = ctx.active_rule
    matching_finding = ctx.matching_finding
    primary_behavioral = ctx.primary_behavioral
    explanation_mode = ctx.explanation_mode
    font_name, font_bold, has_rupee = init_pdf_fonts()
    gen_dt = ctx.generation_date or datetime.now()
    date_str = gen_dt.strftime("%d %B %Y")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName=font_bold,
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0F172A"),
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#0284C7"),
    )
    h1_style = ParagraphStyle(
        "Heading1_Custom",
        parent=styles["Normal"],
        fontName=font_bold,
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#0F172A"),
        spaceAfter=5,
        keepWithNext=True,
    )
    h2_style = ParagraphStyle(
        "Heading2_Custom",
        parent=styles["Normal"],
        fontName=font_bold,
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#1E293B"),
        spaceAfter=3,
        keepWithNext=True,
    )
    body_style = ParagraphStyle(
        "Body_Custom",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#334155"),
        spaceAfter=4,
    )
    callout_style = ParagraphStyle(
        "Callout_Text",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#1E293B"),
    )
    table_cell = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1E293B"),
    )
    table_cell_bold = ParagraphStyle(
        "TableCellBold",
        parent=table_cell,
        fontName=font_bold,
    )
    table_cell_header = ParagraphStyle(
        "TableCellHeader",
        parent=table_cell,
        fontName=font_bold,
        textColor=colors.HexColor("#FFFFFF"),
    )

    story: List[Any] = []

    # -------------------------------------------------------------------------
    # PAGE 1: COVER & EXECUTIVE SUMMARY
    # -------------------------------------------------------------------------
    story.append(Paragraph("<b>NUBRA TRADE REPLAY</b> &nbsp;|&nbsp; POST-TRADE REVIEW SUITE", subtitle_style))
    story.append(Spacer(1, 3))
    story.append(Paragraph("30-Day Trading Behaviour Action Plan", title_style))
    story.append(Spacer(1, 5))
    story.append(
        Paragraph(
            f"<b>Generation Date:</b> {date_str} &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"<b>Review Target:</b> {safe_text(dataset_label)}",
            body_style,
        )
    )
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#0284C7"), spaceAfter=10))

    # Executive Summary Box
    summary_html = (
        "<b>Plan Objective:</b> This document establishes a single, practical habit commitment for the upcoming "
        "30 trading days. The objective is to build execution awareness around high-risk behavioral tendencies "
        "(such as rapid post-loss re-entries, volume surges, or event-day volatility) and maintain consistent "
        "trade pacing. This plan is designed for personal journal tracking and does not constitute investment advice."
    )
    summary_table = Table([[Paragraph(summary_html, body_style)]], colWidths=[487])
    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
                ("PADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.append(summary_table)
    story.append(Spacer(1, 10))

    # Committed Rule Focus Card
    story.append(Paragraph("Selected Habit Rule Commitment", h1_style))
    rule_table_data = [
        [
            Paragraph("<b>Committed Rule:</b>", table_cell_bold),
            Paragraph(f"<b>{safe_text(active_rule['name'])}</b> ({safe_text(active_rule['category'])})", table_cell),
        ],
        [
            Paragraph("<b>Target Pattern:</b>", table_cell_bold),
            Paragraph(safe_text(active_rule["target_pattern"]), table_cell),
        ],
        [
            Paragraph("<b>Trigger Condition:</b>", table_cell_bold),
            Paragraph(safe_text(active_rule.get("trigger", "Specified execution event occurs.")), table_cell),
        ],
        [
            Paragraph("<b>Action to Follow:</b>", table_cell_bold),
            Paragraph(safe_text(active_rule.get("action", active_rule["summary"])), table_cell),
        ],
        [
            Paragraph("<b>Setup Check:</b>", table_cell_bold),
            Paragraph(safe_text(active_rule.get("additional_check", active_rule["rationale"])), table_cell),
        ],
    ]
    rule_table = Table(rule_table_data, colWidths=[120, 367])
    rule_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F0FDF4")),
                ("BOX", (0, 0), (-1, -1), 1.25, colors.HexColor("#16A34A")),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DCFCE7")),
                ("PADDING", (0, 0), (-1, -1), 5),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(rule_table)
    story.append(Spacer(1, 10))

    # Baseline Review Metrics Snapshot
    story.append(Paragraph("Baseline Portfolio Review Metrics", h1_style))
    pf_display = f"{metrics.profit_factor:.2f}" if metrics.profit_factor is not None else "N/A (No losing trades)"

    metrics_table_data = [
        [
            Paragraph("<b>Metric</b>", table_cell_header),
            Paragraph("<b>Baseline Value</b>", table_cell_header),
            Paragraph("<b>Metric</b>", table_cell_header),
            Paragraph("<b>Baseline Value</b>", table_cell_header),
        ],
        [
            Paragraph("Completed Trades", table_cell),
            Paragraph(f"<b>{metrics.total_trades}</b>", table_cell_bold),
            Paragraph("Win Rate (%)", table_cell),
            Paragraph(f"<b>{metrics.win_rate:.1f}%</b> ({metrics.winning_trades}W / {metrics.losing_trades}L)", table_cell_bold),
        ],
        [
            Paragraph("Net Realized P&amp;L", table_cell),
            Paragraph(f"<b>{format_currency_text(metrics.total_net_pnl, has_rupee)}</b>", table_cell_bold),
            Paragraph("Gross Realized P&amp;L", table_cell),
            Paragraph(format_currency_text(metrics.total_gross_pnl, has_rupee), table_cell),
        ],
        [
            Paragraph("Total Charges &amp; Taxes", table_cell),
            Paragraph(format_currency_text(metrics.total_charges, has_rupee).replace("+", ""), table_cell),
            Paragraph("Net Profit Factor", table_cell),
            Paragraph(pf_display, table_cell),
        ],
        [
            Paragraph("Active Trading Days", table_cell),
            Paragraph(str(metrics.active_days), table_cell),
            Paragraph("Avg Net P&amp;L / Trade", table_cell),
            Paragraph(format_currency_text(metrics.avg_net_pnl, has_rupee), table_cell),
        ],
    ]
    metrics_table = Table(metrics_table_data, colWidths=[130, 113, 130, 114])
    metrics_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("PADDING", (0, 0), (-1, -1), 4.5),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story.append(metrics_table)
    story.append(Spacer(1, 10))

    # Disclosures on Page 1
    if is_synthetic:
        synth_box = Table(
            [[
                Paragraph(
                    "<b>⚠️ Synthetic Demonstration Data:</b> This action plan was generated from a synthetic "
                    f"demonstration tradebook (<b>{safe_text(dataset_label)}</b>) designed to demonstrate behavioral "
                    "habit monitoring under controlled conditions. It does not reflect live account trading history.",
                    callout_style,
                )
            ]],
            colWidths=[487],
        )
        synth_box.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FEF3C7")),
                    ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#F59E0B")),
                    ("PADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        story.append(synth_box)
        story.append(Spacer(1, 6))
    else:
        upload_box = Table(
            [[
                Paragraph(
                    f"<b>✓ User Uploaded Tradebook:</b> All figures and evidence in this plan are drawn strictly "
                    f"from your uploaded completed trades CSV (<b>{safe_text(dataset_label)}</b>).",
                    callout_style,
                )
            ]],
            colWidths=[487],
        )
        upload_box.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F0FDF4")),
                    ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#16A34A")),
                    ("PADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        story.append(upload_box)
        story.append(Spacer(1, 6))

    if metrics.is_small_sample:
        sample_box = Table(
            [[
                Paragraph(
                    f"<b>⚠️ Low-Confidence Notice (Small Sample):</b> The reviewed dataset contains only "
                    f"<b>{metrics.total_trades} completed trades</b> (&lt; 10 trades). Statistical identification of "
                    f"recurring habits requires larger trade samples. Interpret baseline numbers as exploratory.",
                    callout_style,
                )
            ]],
            colWidths=[487],
        )
        sample_box.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FEE2E2")),
                    ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#EF4444")),
                    ("PADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        story.append(sample_box)

    # -------------------------------------------------------------------------
    # PAGE 2: WHY THIS RULE WAS SELECTED & SUPPORTING EVIDENCE
    # -------------------------------------------------------------------------
    story.append(PageBreak())

    matching_finding, primary_behavioral, explanation_mode = analyze_review_and_rule(findings, active_rule)

    story.append(Paragraph("Section 2 — Why This Rule Was Selected", h1_style))

    behavioral_findings = [f for f in findings if f.detector_id not in ("clean_baseline", "baseline_clean")]

    if explanation_mode == "matched" and matching_finding is not None:
        clean_summary = clean_currency_text(matching_finding.summary).strip()
        if clean_summary.startswith("Identified "):
            clean_summary = "The review identified " + clean_summary[len("Identified "):]
        if not clean_summary.endswith('.'):
            clean_summary += '.'

        if matching_finding.detector_id == "rapid_reentry" or active_rule.get("id") == "cool_off":
            rule_explanation = (
                f"Committing to <b>{safe_text(active_rule['name'])}</b> directly addresses the observed behavior "
                f"by establishing a structured pause after a losing exit, providing an opportunity to check that "
                f"the next trade is based on an independent setup rather than an impulse re-entry."
            )
        elif matching_finding.detector_id == "high_activity" or active_rule.get("id") == "activity_ceiling":
            rule_explanation = (
                f"Committing to <b>{safe_text(active_rule['name'])}</b> directly addresses the observed high-activity pattern. "
                f"The operational standard is clear: When the fourth completed trade of the day is reached, pause and "
                f"reassess before placing another trade. This protocol emphasizes pausing when the daily trade-count threshold "
                f"is reached, reassessing market conditions and strategy fit, and avoiding unnecessary additional trades."
            )
        elif matching_finding.detector_id in ("expiry_pattern", "expiry_contrast") or active_rule.get("id") == "expiry_sizing":
            rule_explanation = (
                f"Committing to <b>{safe_text(active_rule['name'])}</b> directly addresses the observed performance divergence "
                f"by reducing position size on options expiry days to manage heightened gamma and theta risks."
            )
        else:
            rule_explanation = (
                f"Committing to <b>{safe_text(active_rule['name'])}</b> establishes structured execution habits "
                f"to support disciplined trade selection and risk management."
            )

        why_text = (
            f"This habit rule was directly recommended based on your tradebook review. "
            f"The detector identified <b>{safe_text(matching_finding.title)}</b>. "
            f"{safe_text(clean_summary)} {rule_explanation}"
        )
        other_findings = [f for f in behavioral_findings if f != matching_finding]
        if other_findings:
            other_names = ", ".join(f"{safe_text(f.title)}" for f in other_findings)
            why_text += f"<br/><br/><i>Additional behavioral patterns detected in this review: {other_names}.</i>"
    elif explanation_mode == "user_custom" and primary_behavioral is not None:
        why_text = (
            f"The primary behavioral pattern detected in your review was <b>{safe_text(primary_behavioral.title)}</b> "
            f"({safe_text(primary_behavioral.metric_label)}: {safe_text(primary_behavioral.metric_value)}). "
            f"However, you have chosen to commit to <b>{safe_text(active_rule['name'])}</b> ({safe_text(active_rule['category'])}) "
            f"as your personal focus for the next 30 days. The supporting evidence below summarizes the observed pattern in your tradebook."
        )
    else:
        why_text = (
            f"<b>Clean Baseline Review:</b> No configured behavioral patterns met detection criteria during this review period. "
            f"Across the {metrics.total_trades} completed trades analyzed, trade spacing, daily trade counts, and event-day distributions "
            f"remained within normal baseline boundaries. You selected <b>{safe_text(active_rule['name'])}</b> as a proactive discipline "
            f"standard to reinforce execution consistency and maintain disciplined pacing over the next 30 days."
        )

    story.append(Paragraph(why_text, body_style))
    story.append(Spacer(1, 6))

    # Supporting Evidence Table (if a behavioral pattern was detected)
    active_evidence_finding = matching_finding or primary_behavioral
    if active_evidence_finding is not None and active_evidence_finding.detector_id not in ("clean_baseline", "baseline_clean"):
        story.append(Paragraph(f"Supporting Evidence: {safe_text(active_evidence_finding.title)}", h2_style))

        if active_evidence_finding.detector_id == "rapid_reentry":
            ev_df = active_evidence_finding.evidence_df
            if isinstance(ev_df, pd.DataFrame) and not ev_df.empty:
                max_display = min(6, len(ev_df))
                story.append(
                    Paragraph(
                        f"<b>Total Flagged Events: {len(ev_df)}</b> (showing first {max_display} supporting events below):",
                        callout_style,
                    )
                )
                story.append(Spacer(1, 2))
                ev_data = [
                    [
                        Paragraph("<b>Flagged Trade ID</b>", table_cell_header),
                        Paragraph("<b>Preceding Loss Trade</b>", table_cell_header),
                        Paragraph("<b>Next Entry Time</b>", table_cell_header),
                        Paragraph("<b>Time Gap</b>", table_cell_header),
                        Paragraph("<b>Preceding Loss Net P&amp;L</b>", table_cell_header),
                        Paragraph("<b>Flagged Trade Net P&amp;L</b>", table_cell_header),
                    ]
                ]
                for r_idx in range(max_display):
                    row = ev_df.iloc[r_idx]
                    prev_pnl = float(row.get("preceding_net_pnl", 0.0))
                    curr_pnl = float(row.get("net_pnl", 0.0))
                    gap_val = float(row.get("gap_minutes", 0.0))
                    ev_data.append([
                        Paragraph(f"{row['date']}<br/><b>{row['trade_id']}</b>", table_cell),
                        Paragraph(f"<b>{row.get('preceding_trade_id', '')}</b><br/>Exit: {row.get('preceding_exit_time', '')}", table_cell),
                        Paragraph(str(row.get("entry_time", "")), table_cell),
                        Paragraph(f"<b>{gap_val:.1f} min</b>", table_cell_bold),
                        Paragraph(format_currency_text(prev_pnl, has_rupee), table_cell),
                        Paragraph(format_currency_text(curr_pnl, has_rupee), table_cell),
                    ])

                ev_table = Table(ev_data, colWidths=[90, 85, 75, 55, 92, 90])
                ev_table.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                            ("PADDING", (0, 0), (-1, -1), 4),
                            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
                            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ]
                    )
                )
                story.append(ev_table)
                if len(ev_df) > max_display:
                    story.append(Spacer(1, 2))
                    story.append(Paragraph(f"<i>Showing first {max_display} of {len(ev_df)} total flagged rapid re-entry events.</i>", callout_style))

                caption_html = (
                    "<b>Note on P&amp;L Attribution:</b> Preceding Loss Net P&amp;L reflects the loss that immediately preceded "
                    "the rapid entry. Flagged Trade Net P&amp;L reflects the outcome of the subsequent trade. These are distinct events; "
                    "observing a short time gap does not mean the gap caused the subsequent trade result."
                )
                story.append(Spacer(1, 4))
                story.append(Paragraph(caption_html, callout_style))

        elif active_evidence_finding.detector_id == "high_activity":
            ev_df = active_evidence_finding.evidence_df
            if isinstance(ev_df, pd.DataFrame) and not ev_df.empty:
                ev_data = [
                    [
                        Paragraph("<b>Trading Date</b>", table_cell_header),
                        Paragraph("<b>Daily Trades Executed</b>", table_cell_header),
                        Paragraph("<b>Baseline Daily Median</b>", table_cell_header),
                        Paragraph("<b>Activity Threshold</b>", table_cell_header),
                        Paragraph("<b>Daily Net Realized P&amp;L</b>", table_cell_header),
                    ]
                ]
                for _, row in ev_df.iterrows():
                    ev_data.append([
                        Paragraph(str(row.get("Date", "")), table_cell_bold),
                        Paragraph(f"<b>{row.get('Trade Count', '')} trades</b>", table_cell_bold),
                        Paragraph(f"{row.get('Baseline Median', '')} trades/day", table_cell),
                        Paragraph(f"{row.get('Threshold', '')} trades/day", table_cell),
                        Paragraph(clean_currency_text(str(row.get("Daily Net P&L", ""))), table_cell),
                    ])
                ev_table = Table(ev_data, colWidths=[97, 100, 95, 95, 100])
                ev_table.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                            ("PADDING", (0, 0), (-1, -1), 4),
                            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
                            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ]
                    )
                )
                story.append(ev_table)
                story.append(Spacer(1, 4))
                story.append(
                    Paragraph(
                        "<b>Note:</b> High trade activity on losing days is an observed association with difficult market conditions "
                        "or overtrading. It does not prove that trade count alone caused the financial loss.",
                        callout_style,
                    )
                )

        elif active_evidence_finding.detector_id in ("expiry_pattern", "expiry_contrast"):
            ev_df = active_evidence_finding.evidence_df
            if isinstance(ev_df, pd.DataFrame) and not ev_df.empty:
                ev_data = [
                    [
                        Paragraph("<b>Session Group</b>", table_cell_header),
                        Paragraph("<b>Completed Trades</b>", table_cell_header),
                        Paragraph("<b>Total Group Net P&amp;L</b>", table_cell_header),
                        Paragraph("<b>Avg Net P&amp;L / Trade</b>", table_cell_header),
                    ]
                ]
                for _, row in ev_df.iterrows():
                    ev_data.append([
                        Paragraph(str(row.get("Group", "")), table_cell_bold),
                        Paragraph(str(row.get("Trade Count", "")), table_cell),
                        Paragraph(clean_currency_text(str(row.get("Total Net P&L", ""))), table_cell),
                        Paragraph(clean_currency_text(str(row.get("Avg Net P&L / Trade", ""))), table_cell_bold),
                    ])
                ev_table = Table(ev_data, colWidths=[140, 100, 120, 127])
                ev_table.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                            ("PADDING", (0, 0), (-1, -1), 4),
                            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
                            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ]
                    )
                )
                story.append(ev_table)
                story.append(Spacer(1, 4))
                story.append(
                    Paragraph(
                        "<b>Note:</b> Options expiration involves elevated gamma and theta decay. Performance contrast reflects historical "
                        "group averages and does not prove that the calendar day caused individual trade losses.",
                        callout_style,
                    )
                )

    story.append(Spacer(1, 8))

    # Guardrail on Association vs. Causation
    guardrail_html = (
        "<b>Understanding Observed Patterns (Association vs. Causation):</b> "
        "The patterns identified by Nubra Trade Replay describe historical statistical associations in your completed tradebook. "
        "They do not prove that entering quickly or trading on expiry caused losses, nor does following this plan guarantee future trading profits. "
        "The sole purpose of habit tracking is to improve execution self-awareness."
    )
    guardrail_table = Table([[Paragraph(guardrail_html, callout_style)]], colWidths=[487])
    guardrail_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EFF6FF")),
                ("BOX", (0, 0), (-1, -1), 0.75, colors.HexColor("#3B82F6")),
                ("PADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(guardrail_table)
    story.append(Spacer(1, 10))

    # Section 3: 30-Day Action Checklist
    story.append(Paragraph("Section 3 — The 30-Day Action Checklist", h1_style))
    story.append(Paragraph("Follow this daily execution protocol during every live trading session over the next 30 days:", body_style))
    story.append(Spacer(1, 3))

    daily_items = active_rule.get("checklist", [])
    checklist_rows = []
    for item in daily_items:
        checklist_rows.append([
            Paragraph("<b>[ &nbsp; ]</b>", table_cell_bold),
            Paragraph(safe_text(item), table_cell),
        ])

    chk_table = Table(checklist_rows, colWidths=[35, 452])
    chk_table.setStyle(
        TableStyle(
            [
                ("PADDING", (0, 0), (-1, -1), 3.5),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ]
        )
    )
    story.append(chk_table)
    story.append(Spacer(1, 8))

    # Weekly Milestones
    story.append(Paragraph("Weekly Review Milestones", h2_style))
    if active_rule.get("id") == "activity_ceiling":
        weekly_data = [
            [
                Paragraph("<b>Milestone</b>", table_cell_header),
                Paragraph("<b>Weekly Objective &amp; Adherence Audit</b>", table_cell_header),
                Paragraph("<b>Status</b>", table_cell_header),
            ],
            [
                Paragraph("<b>Week 1</b><br/>(Days 1–7)", table_cell),
                Paragraph(
                    f"<b>Setup & Calibration:</b> Log daily trade counts in real time. Confirm strict adherence to "
                    f"<b>{safe_text(active_rule['name'])}</b>: When the fourth completed trade of the day is reached, pause and reassess before placing another trade.",
                    table_cell,
                ),
                Paragraph("[ &nbsp; ] Complete", table_cell),
            ],
            [
                Paragraph("<b>Week 2</b><br/>(Days 8–14)", table_cell),
                Paragraph(
                    f"<b>Mid-Month Audit:</b> Calculate your interim adherence rate for the first two weeks. Verify that qualifying "
                    f"trigger events (<i>When the fourth completed trade of the day is reached, pause and reassess before placing another trade.</i>) were handled strictly according to the rule.",
                    table_cell,
                ),
                Paragraph("[ &nbsp; ] Complete", table_cell),
            ],
            [
                Paragraph("<b>Week 3</b><br/>(Days 15–21)", table_cell),
                Paragraph(
                    f"<b>Handling Pressure & High Activity:</b> Review trading during fast-moving or choppy sessions. "
                    f"Confirm you paused when reaching the fourth trade, reassessed market conditions and strategy fit, and avoided unnecessary additional trades.",
                    table_cell,
                ),
                Paragraph("[ &nbsp; ] Complete", table_cell),
            ],
            [
                Paragraph("<b>Week 4</b><br/>(Days 22–30)", table_cell),
                Paragraph(
                    "<b>Month-End Consolidation:</b> Export your completed 30-day tradebook and replay it in Nubra Trade Replay "
                    "to evaluate trade frequency consistency and compare pattern recurrence against your pre-commitment baseline.",
                    table_cell,
                ),
                Paragraph("[ &nbsp; ] Complete", table_cell),
            ],
        ]
    else:
        clean_action = str(active_rule.get('action', active_rule['summary'])).rstrip('.')
        clean_trigger = str(active_rule.get('trigger', '')).rstrip('.')
        weekly_data = [
            [
                Paragraph("<b>Milestone</b>", table_cell_header),
                Paragraph("<b>Weekly Objective &amp; Adherence Audit</b>", table_cell_header),
                Paragraph("<b>Status</b>", table_cell_header),
            ],
            [
                Paragraph("<b>Week 1</b><br/>(Days 1–7)", table_cell),
                Paragraph(
                    f"<b>Setup & Calibration:</b> Log every trade immediately after exit. Confirm strict adherence to "
                    f"<b>{safe_text(active_rule['name'])}</b>: {safe_text(clean_action)}.",
                    table_cell,
                ),
                Paragraph("[ &nbsp; ] Complete", table_cell),
            ],
            [
                Paragraph("<b>Week 2</b><br/>(Days 8–14)", table_cell),
                Paragraph(
                    f"<b>Mid-Month Audit:</b> Calculate your interim adherence rate for the first two weeks. Verify that qualifying "
                    f"trigger events (<i>{safe_text(clean_trigger)}</i>) were handled strictly according to the rule.",
                    table_cell,
                ),
                Paragraph("[ &nbsp; ] Complete", table_cell),
            ],
            [
                Paragraph("<b>Week 3</b><br/>(Days 15–21)", table_cell),
                Paragraph(
                    f"<b>Handling Pressure & Drawdowns:</b> Review trading during choppy sessions or losing streaks. "
                    f"Did you consistently maintain your commitment to <b>{safe_text(active_rule['name'])}</b> without deviation?",
                    table_cell,
                ),
                Paragraph("[ &nbsp; ] Complete", table_cell),
            ],
            [
                Paragraph("<b>Week 4</b><br/>(Days 22–30)", table_cell),
                Paragraph(
                    "<b>Month-End Consolidation:</b> Export your completed 30-day tradebook and replay it in Nubra Trade Replay "
                    "to evaluate habit consistency and compare pattern recurrence against your pre-commitment baseline.",
                    table_cell,
                ),
                Paragraph("[ &nbsp; ] Complete", table_cell),
            ],
        ]
    weekly_table = Table(weekly_data, colWidths=[75, 342, 70])
    weekly_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("PADDING", (0, 0), (-1, -1), 4),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(weekly_table)

    # -------------------------------------------------------------------------
    # PAGE 3: MEASURABLE SUCCESS TRACKING & 30-DAY LOG
    # -------------------------------------------------------------------------
    story.append(PageBreak())

    story.append(Paragraph("Section 4 — Measurable Success Tracking", h1_style))
    story.append(
        Paragraph(
            "Track rule adherence as a behavioral discipline measure rather than a financial profit guarantee. "
            "Record whether the rule was followed whenever the qualifying trigger condition occurs.",
            body_style,
        )
    )
    story.append(Spacer(1, 4))

    # Definition & Formula Table
    formula_html = (
        "<b>Adherence Rate Formula:</b><br/>"
        "&nbsp;&nbsp;<b>Adherence Rate (%) = (Trigger Events Where Rule Was Followed ÷ Total Qualifying Trigger Events) × 100</b><br/>"
        "<i>Note: If no qualifying trigger events occur during the month, adherence rate is reported as: "
        "<b>'Not applicable — no qualifying events occurred'</b> (do not divide by zero or assume 100%).</i>"
    )
    formula_table = Table([[Paragraph(formula_html, body_style)]], colWidths=[487])
    formula_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
                ("PADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(formula_table)
    story.append(Spacer(1, 6))

    # Core Metrics Tracking Definitions
    metrics_tracking_data = [
        [
            Paragraph("<b>Tracking Measure</b>", table_cell_header),
            Paragraph("<b>Operational Definition</b>", table_cell_header),
            Paragraph("<b>Measurement Method</b>", table_cell_header),
        ],
        [
            Paragraph("<b>Number of Qualifying Trigger Events</b>", table_cell_bold),
            Paragraph(f"Total count of trading instances where the trigger condition occurs: <i>{safe_text(active_rule.get('trigger', ''))}</i>", table_cell),
            Paragraph("Recorded in daily tracker as 'Trigger Occurred? = Yes'", table_cell),
        ],
        [
            Paragraph("<b>Number of Events Where Rule Followed</b>", table_cell_bold),
            Paragraph(f"Count of qualifying trigger instances where the committed rule action was strictly adhered to: <i>{safe_text(active_rule.get('action', ''))}</i>", table_cell),
            Paragraph("Recorded in daily tracker as 'Rule Followed? = Yes'", table_cell),
        ],
        [
            Paragraph("<b>Number of Violations</b>", table_cell_bold),
            Paragraph("Count of qualifying trigger instances where the committed rule was breached or not followed", table_cell),
            Paragraph("Calculated as: Qualifying Trigger Events − Events Where Rule Followed (logged as 'Rule Followed? = No')", table_cell),
        ],
        [
            Paragraph("<b>Rule Adherence Rate</b>", table_cell_bold),
            Paragraph("Calculated adherence percentage using the formula above: (Events Followed ÷ Total Qualifying Events) × 100", table_cell),
            Paragraph("Calculated at month-end. If 0 qualifying trigger events occurred, reported as: 'Not applicable — no qualifying events occurred'", table_cell),
        ],
        [
            Paragraph("<b>Monitoring Period Review</b>", table_cell_bold),
            Paragraph("Separate review of next-month tradebook; evaluated independently without mixing future results into historical baseline", table_cell),
            Paragraph("Re-evaluate tradebook at month-end. No financial profit or loss reduction is promised or guaranteed.", table_cell),
        ],
    ]
    track_def_table = Table(metrics_tracking_data, colWidths=[125, 205, 157])
    track_def_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("PADDING", (0, 0), (-1, -1), 4),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story.append(track_def_table)
    story.append(Spacer(1, 8))

    # Complete 30-Day Adherence Log (All 30 Days)
    story.append(Paragraph("Complete 30-Day Adherence Log (Daily Journal Worksheets)", h2_style))
    story.append(
        Paragraph(
            "Fill in each trading session below. If no qualifying trigger event occurred on that day, mark 'N/A'.",
            callout_style,
        )
    )
    story.append(Spacer(1, 3))

    log_data = [
        [
            Paragraph("<b>Day / Date</b>", table_cell_header),
            Paragraph("<b>Trigger Occurred?</b>", table_cell_header),
            Paragraph("<b>Rule Followed?</b>", table_cell_header),
            Paragraph("<b>Trader Notes &amp; Emotional State</b>", table_cell_header),
        ]
    ]

    for day_num in range(1, 31):
        log_data.append([
            Paragraph(f"Day {day_num} / Day {day_num:02d}: ___/___", table_cell),
            Paragraph("[ ] Yes &nbsp; [ ] No", table_cell),
            Paragraph("[ ] Yes &nbsp; [ ] No &nbsp; [ ] N/A", table_cell),
            Paragraph(" ", table_cell),
        ])

    # Table with repeatRows=1 so header repeats automatically when breaking across pages
    log_table = Table(log_data, colWidths=[90, 105, 105, 187], repeatRows=1)
    log_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("PADDING", (0, 0), (-1, -1), 2.5),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story.append(log_table)
    story.append(Spacer(1, 10))

    # Day-30 Reflection Section
    story.append(Paragraph("Day-30 Final Reflection", h2_style))
    reflection_html = (
        "<b>1. Final Adherence Rate:</b> Qualifying events followed: [ ______ ] ÷ Total trigger events: [ ______ ] = <b>[ &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; ]%</b> "
        "<i>(or [ ] Not applicable — no qualifying events occurred)</i><br/>"
        "<b>2. Key Challenge:</b> Under what market conditions or emotional state was it most difficult to follow this rule? "
        "__________________________________________________________________________________<br/>"
        "<b>3. Next Month Commitment:</b> What habit rule or risk boundary will you carry forward into next month? "
        "__________________________________________________________________________________"
    )
    refl_table = Table([[Paragraph(reflection_html, body_style)]], colWidths=[487])
    refl_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
                ("PADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.append(refl_table)
    story.append(Spacer(1, 10))

    # Section 5: Limitations & Disclaimers
    story.append(Paragraph("Section 5 — Important Limitations &amp; Disclaimers", h1_style))
    disclaimer_html = (
        "• <b>Educational Self-Monitoring Aid:</b> Nubra Trade Replay is a post-trade review prototype designed solely to assist "
        "in personal journaling and execution self-awareness. It does not provide financial advice, trading signals, or buy/sell recommendations.<br/>"
        "• <b>No Guarantee of Financial Outcomes:</b> Past behavioral habits and trade correlations do not predict future market results. "
        "Adhering to habit rules cannot guarantee trading profits or protect against market drawdowns.<br/>"
        "• <b>Dataset Attribution:</b> All figures and findings are derived directly from the currently selected dataset without external "
        "extrapolation. If this dataset is a synthetic demonstration tradebook, it is intended solely for illustrative habit monitoring."
    )
    disc_table = Table([[Paragraph(disclaimer_html, callout_style)]], colWidths=[487])
    disc_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F1F5F9")),
                ("BOX", (0, 0), (-1, -1), 0.75, colors.HexColor("#94A3B8")),
                ("PADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(disc_table)

    # Build document
    doc.build(story, canvasmaker=NumberedCanvas)
    return buffer.getvalue()


# -----------------------------------------------------------------------------
# SECONDARY EXPORT: CANONICAL DOCX GENERATION
# -----------------------------------------------------------------------------

def generate_action_plan_docx(
    dataset_label: Optional[Union[ReportContext, str]] = None,
    is_synthetic: Optional[bool] = None,
    metrics: Optional[SummaryMetrics] = None,
    findings: Optional[List[Finding]] = None,
    active_rule: Optional[Dict[str, Any]] = None,
    generation_date: Optional[datetime] = None,
    df_trades: Optional[pd.DataFrame] = None,
    context: Optional[ReportContext] = None,
) -> bytes:
    """
    Canonical entry point for editable Microsoft Word (.docx) 30-Day Action Plan generation.
    Accepts a normalized ReportContext or individual parameters.
    """
    if isinstance(dataset_label, ReportContext):
        ctx = dataset_label
    elif context is not None:
        ctx = context
    else:
        ctx = build_report_context(
            dataset_label=str(dataset_label or ""),
            is_synthetic=bool(is_synthetic),
            metrics=metrics,
            findings=findings or [],
            active_rule=active_rule or {},
            df_trades=df_trades,
            generation_date=generation_date,
        )

    dataset_label = ctx.dataset_label
    is_synthetic = ctx.is_synthetic
    metrics = ctx.metrics
    findings = ctx.findings
    active_rule = ctx.active_rule
    matching_finding = ctx.matching_finding
    primary_behavioral = ctx.primary_behavioral
    explanation_mode = ctx.explanation_mode
    gen_dt = ctx.generation_date or datetime.now()
    date_str = gen_dt.strftime("%d %B %Y")
    pf_display = f"{metrics.profit_factor:.2f}" if metrics.profit_factor is not None else "N/A (No losing trades)"

    doc = docx.Document()

    # Set standard margins (0.75 in)
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    # Header branding
    p_header = doc.add_paragraph()
    run_brand = p_header.add_run("NUBRA TRADE REPLAY | POST-TRADE REVIEW SUITE")
    run_brand.font.name = "Arial"
    run_brand.font.size = Pt(9)
    run_brand.font.color.rgb = RGBColor(2, 132, 199)
    run_brand.bold = True

    # Title
    p_title = doc.add_paragraph()
    run_title = p_title.add_run("30-Day Trading Behaviour Action Plan")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(18)
    run_title.font.color.rgb = RGBColor(15, 23, 42)
    run_title.bold = True

    p_meta = doc.add_paragraph()
    p_meta.add_run(f"Generation Date: {date_str}   |   Review Target: {dataset_label}").italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Executive Summary
    p_sum = doc.add_paragraph()
    p_sum_bold = p_sum.add_run("Plan Objective: ")
    p_sum_bold.bold = True
    p_sum.add_run(
        "This document establishes a single, practical habit commitment for the upcoming 30 trading days. "
        "The objective is to cultivate execution self-awareness around behavioral tendencies (such as rapid "
        "post-loss entries or volume surges) and maintain consistent pacing. This plan is designed for personal "
        "journal tracking and does not constitute financial advice."
    )

    # Committed Rule Focus
    doc.add_heading("Selected Habit Rule Commitment", level=1)
    t_rule = doc.add_table(rows=5, cols=2)
    t_rule.alignment = WD_TABLE_ALIGNMENT.CENTER
    rule_data = [
        ("Committed Rule:", f"{active_rule['name']} ({active_rule['category']})"),
        ("Target Pattern:", active_rule["target_pattern"]),
        ("Trigger Condition:", active_rule.get("trigger", "Specified execution event occurs.")),
        ("Action to Follow:", active_rule.get("action", active_rule["summary"])),
        ("Setup Check:", active_rule.get("additional_check", active_rule["rationale"])),
    ]
    for row_idx, (label, val) in enumerate(rule_data):
        row = t_rule.rows[row_idx]
        row.cells[0].paragraphs[0].add_run(label).bold = True
        row.cells[1].paragraphs[0].add_run(val)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Baseline Portfolio Review Metrics Snapshot
    doc.add_heading("Baseline Portfolio Review Metrics", level=1)
    t_metrics = doc.add_table(rows=5, cols=4)
    t_metrics.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Metric", "Baseline Value", "Metric", "Baseline Value"]
    for i, h in enumerate(headers):
        t_metrics.rows[0].cells[i].paragraphs[0].add_run(h).bold = True

    m_data = [
        ("Completed Trades", str(metrics.total_trades), "Win Rate (%)", f"{metrics.win_rate:.1f}% ({metrics.winning_trades}W / {metrics.losing_trades}L)"),
        ("Net Realized P&L", format_currency_text(metrics.total_net_pnl, True), "Gross Realized P&L", format_currency_text(metrics.total_gross_pnl, True)),
        ("Total Charges & Taxes", f"₹{metrics.total_charges:,.2f}", "Net Profit Factor", pf_display),
        ("Active Trading Days", str(metrics.active_days), "Avg Net P&L / Trade", format_currency_text(metrics.avg_net_pnl, True)),
    ]
    for row_idx, (m1, v1, m2, v2) in enumerate(m_data, start=1):
        row = t_metrics.rows[row_idx]
        row.cells[0].paragraphs[0].add_run(m1)
        row.cells[1].paragraphs[0].add_run(v1).bold = True
        row.cells[2].paragraphs[0].add_run(m2)
        row.cells[3].paragraphs[0].add_run(v2).bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Notices
    if is_synthetic:
        p_syn = doc.add_paragraph()
        run_syn = p_syn.add_run(
            f"⚠️ Synthetic Demonstration Data: This action plan was generated from a synthetic "
            f"demonstration tradebook ({dataset_label}) for habit-monitoring and educational purposes. "
            f"It does not reflect live account trading history."
        )
        run_syn.font.color.rgb = RGBColor(180, 83, 9)
        run_syn.bold = True
    else:
        p_up = doc.add_paragraph()
        run_up = p_up.add_run(
            f"✓ User Uploaded Tradebook: All figures and evidence in this plan are drawn strictly "
            f"from your uploaded completed trades CSV ({dataset_label})."
        )
        run_up.font.color.rgb = RGBColor(22, 163, 74)
        run_up.bold = True

    if metrics.is_small_sample:
        p_sam = doc.add_paragraph()
        run_sam = p_sam.add_run(
            f"⚠️ Low-Confidence Notice (Small Sample): The reviewed dataset contains only {metrics.total_trades} completed trades "
            f"(< 10 trades). Statistical identification of recurring habits requires larger trade samples."
        )
        run_sam.font.color.rgb = RGBColor(220, 38, 38)
        run_sam.bold = True

    # Section 2: Why This Rule Was Selected
    doc.add_page_break()
    doc.add_heading("Section 2 — Why This Rule Was Selected", level=1)

    matching_finding, primary_behavioral, explanation_mode = analyze_review_and_rule(findings, active_rule)
    behavioral_findings = [f for f in findings if f.detector_id not in ("clean_baseline", "baseline_clean")]

    p_why = doc.add_paragraph()
    if explanation_mode == "matched" and matching_finding is not None:
        clean_summary = clean_currency_text(matching_finding.summary).strip()
        if clean_summary.startswith("Identified "):
            clean_summary = "The review identified " + clean_summary[len("Identified "):]
        if not clean_summary.endswith('.'):
            clean_summary += '.'

        if matching_finding.detector_id == "rapid_reentry" or active_rule.get("id") == "cool_off":
            rule_explanation = (
                f"Committing to {active_rule['name']} directly addresses the observed behavior "
                f"by establishing a structured pause after a losing exit, providing an opportunity to check that "
                f"the next trade is based on an independent setup rather than an impulse re-entry."
            )
        elif matching_finding.detector_id == "high_activity" or active_rule.get("id") == "activity_ceiling":
            rule_explanation = (
                f"Committing to {active_rule['name']} directly addresses the observed high-activity pattern. "
                f"The operational standard is clear: When the fourth completed trade of the day is reached, pause and "
                f"reassess before placing another trade. This protocol emphasizes pausing when the daily trade-count threshold "
                f"is reached, reassessing market conditions and strategy fit, and avoiding unnecessary additional trades."
            )
        elif matching_finding.detector_id in ("expiry_pattern", "expiry_contrast") or active_rule.get("id") == "expiry_sizing":
            rule_explanation = (
                f"Committing to {active_rule['name']} directly addresses the observed performance divergence "
                f"by reducing position size on options expiry days to manage heightened gamma and theta risks."
            )
        else:
            rule_explanation = (
                f"Committing to {active_rule['name']} establishes structured execution habits "
                f"to support disciplined trade selection and risk management."
            )

        p_why.add_run(
            f"This habit rule was directly recommended based on your tradebook review. "
            f"The detector identified {matching_finding.title}. "
            f"{clean_summary} {rule_explanation}"
        )
        other_findings = [f for f in behavioral_findings if f != matching_finding]
        if other_findings:
            p_other = doc.add_paragraph()
            other_names = ", ".join(f.title for f in other_findings)
            p_other.add_run(f"Additional behavioral patterns detected in this review: {other_names}.").italic = True
    elif explanation_mode == "user_custom" and primary_behavioral is not None:
        p_why.add_run(
            f"The primary behavioral pattern detected in your review was {primary_behavioral.title} "
            f"({primary_behavioral.metric_label}: {primary_behavioral.metric_value}). "
            f"However, you have chosen to commit to {active_rule['name']} ({active_rule['category']}) "
            f"as your personal focus for the next 30 days. The supporting evidence below summarizes the observed pattern."
        )
    else:
        p_why.add_run(
            f"Clean Baseline Review: No configured behavioral patterns met detection criteria during this review period. "
            f"Across the {metrics.total_trades} completed trades analyzed, trade spacing, daily trade counts, and event-day distributions "
            f"remained within normal baseline boundaries. You selected {active_rule['name']} as a proactive discipline "
            f"standard to reinforce execution consistency and maintain disciplined pacing over the next 30 days."
        )

    # Evidence Table in Word
    active_ev = matching_finding or primary_behavioral
    if active_ev is not None and active_ev.detector_id not in ("clean_baseline", "baseline_clean"):
        doc.add_heading(f"Supporting Evidence: {active_ev.title}", level=2)

        if active_ev.detector_id == "rapid_reentry":
            ev_df = active_ev.evidence_df
            if isinstance(ev_df, pd.DataFrame) and not ev_df.empty:
                max_display = min(6, len(ev_df))
                p_cnt = doc.add_paragraph()
                p_cnt.add_run(f"Total Flagged Events: {len(ev_df)} (showing first {max_display} supporting events below):").bold = True

                t_ev = doc.add_table(rows=max_display + 1, cols=6)
                t_ev.alignment = WD_TABLE_ALIGNMENT.CENTER
                ev_headers = ["Flagged Trade ID", "Preceding Loss Trade", "Next Entry Time", "Time Gap", "Preceding Loss Net P&L", "Flagged Trade Net P&L"]
                for i, h in enumerate(ev_headers):
                    t_ev.rows[0].cells[i].paragraphs[0].add_run(h).bold = True

                for r_idx in range(max_display):
                    row = ev_df.iloc[r_idx]
                    prev_pnl = float(row.get("preceding_net_pnl", 0.0))
                    curr_pnl = float(row.get("net_pnl", 0.0))
                    gap_val = float(row.get("gap_minutes", 0.0))
                    r_cells = t_ev.rows[r_idx + 1].cells
                    r_cells[0].paragraphs[0].add_run(f"{row['date']}\n{row['trade_id']}")
                    r_cells[1].paragraphs[0].add_run(f"{row.get('preceding_trade_id', '')}\nExit: {row.get('preceding_exit_time', '')}")
                    r_cells[2].paragraphs[0].add_run(str(row.get("entry_time", "")))
                    r_cells[3].paragraphs[0].add_run(f"{gap_val:.1f} min").bold = True
                    r_cells[4].paragraphs[0].add_run(format_currency_text(prev_pnl, True))
                    r_cells[5].paragraphs[0].add_run(format_currency_text(curr_pnl, True))

                if len(ev_df) > max_display:
                    p_more = doc.add_paragraph()
                    p_more.add_run(f"Showing first {max_display} of {len(ev_df)} total flagged rapid re-entry events.").italic = True

                p_cap = doc.add_paragraph()
                p_cap.add_run(
                    "Note on P&L Attribution: Preceding Loss Net P&L reflects the loss that immediately preceded "
                    "the rapid entry. Flagged Trade Net P&L reflects the outcome of the subsequent trade. These are distinct events; "
                    "observing a short time gap does not mean the gap caused the subsequent trade result."
                ).italic = True

        elif active_ev.detector_id == "high_activity":
            ev_df = active_ev.evidence_df
            if isinstance(ev_df, pd.DataFrame) and not ev_df.empty:
                t_ev = doc.add_table(rows=len(ev_df) + 1, cols=5)
                t_ev.alignment = WD_TABLE_ALIGNMENT.CENTER
                ev_headers = ["Trading Date", "Daily Trades Executed", "Baseline Daily Median", "Activity Threshold", "Daily Net Realized P&L"]
                for i, h in enumerate(ev_headers):
                    t_ev.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
                for r_idx, (_, row) in enumerate(ev_df.iterrows(), start=1):
                    r_cells = t_ev.rows[r_idx].cells
                    r_cells[0].paragraphs[0].add_run(str(row.get("Date", ""))).bold = True
                    r_cells[1].paragraphs[0].add_run(f"{row.get('Trade Count', '')} trades")
                    r_cells[2].paragraphs[0].add_run(f"{row.get('Baseline Median', '')} trades/day")
                    r_cells[3].paragraphs[0].add_run(f"{row.get('Threshold', '')} trades/day")
                    r_cells[4].paragraphs[0].add_run(clean_currency_text(str(row.get("Daily Net P&L", ""))))

        elif active_ev.detector_id in ("expiry_pattern", "expiry_contrast"):
            ev_df = active_ev.evidence_df
            if isinstance(ev_df, pd.DataFrame) and not ev_df.empty:
                t_ev = doc.add_table(rows=len(ev_df) + 1, cols=4)
                t_ev.alignment = WD_TABLE_ALIGNMENT.CENTER
                ev_headers = ["Session Group", "Completed Trades", "Total Group Net P&L", "Avg Net P&L / Trade"]
                for i, h in enumerate(ev_headers):
                    t_ev.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
                for r_idx, (_, row) in enumerate(ev_df.iterrows(), start=1):
                    r_cells = t_ev.rows[r_idx].cells
                    r_cells[0].paragraphs[0].add_run(str(row.get("Group", ""))).bold = True
                    r_cells[1].paragraphs[0].add_run(str(row.get("Trade Count", "")))
                    r_cells[2].paragraphs[0].add_run(clean_currency_text(str(row.get("Total Net P&L", ""))))
                    r_cells[3].paragraphs[0].add_run(clean_currency_text(str(row.get("Avg Net P&L / Trade", ""))))

    # Causation guardrail
    p_cg = doc.add_paragraph()
    p_cg_bold = p_cg.add_run("Understanding Observed Patterns (Association vs. Causation): ")
    p_cg_bold.bold = True
    p_cg.add_run(
        "The patterns identified above represent historical associations in your completed tradebook. "
        "They do not prove that entering quickly or trading on expiry caused losses, nor does following this plan "
        "guarantee future trading profits. The goal is solely to build execution self-awareness."
    )

    # Section 3: 30-Day Action Checklist
    doc.add_heading("Section 3 — The 30-Day Action Checklist", level=1)
    doc.add_paragraph("Follow this daily execution protocol during every live trading session over the next 30 days:")
    for item in active_rule.get("checklist", []):
        p_chk = doc.add_paragraph()
        p_chk.add_run("[   ]  ").bold = True
        p_chk.add_run(item)

    doc.add_heading("Weekly Review Milestones", level=2)
    t_week = doc.add_table(rows=5, cols=3)
    t_week.rows[0].cells[0].paragraphs[0].add_run("Milestone").bold = True
    t_week.rows[0].cells[1].paragraphs[0].add_run("Weekly Objective & Adherence Audit").bold = True
    t_week.rows[0].cells[2].paragraphs[0].add_run("Status").bold = True

    if active_rule.get("id") == "activity_ceiling":
        weeks = [
            (
                "Week 1 (Days 1–7)",
                f"Setup & Calibration: Log daily trade counts in real time. Confirm strict adherence to {active_rule['name']}: When the fourth completed trade of the day is reached, pause and reassess before placing another trade.",
                "[   ] Complete",
            ),
            (
                "Week 2 (Days 8–14)",
                "Mid-Month Audit: Calculate interim adherence rate for first two weeks. Verify that qualifying trigger events ('When the fourth completed trade of the day is reached, pause and reassess before placing another trade.') were handled strictly according to the rule.",
                "[   ] Complete",
            ),
            (
                "Week 3 (Days 15–21)",
                "Handling Pressure & High Activity: Review trading during fast-moving or choppy sessions. Confirm you paused when reaching the fourth trade, reassessed market conditions and strategy fit, and avoided unnecessary additional trades.",
                "[   ] Complete",
            ),
            (
                "Week 4 (Days 22–30)",
                "Month-End Consolidation: Export completed 30-day tradebook and replay in Nubra Trade Replay to evaluate trade frequency consistency and compare pattern recurrence against baseline.",
                "[   ] Complete",
            ),
        ]
    else:
        clean_action = str(active_rule.get('action', active_rule['summary'])).rstrip('.')
        clean_trigger = str(active_rule.get('trigger', '')).rstrip('.')
        weeks = [
            (
                "Week 1 (Days 1–7)",
                f"Setup & Calibration: Log every trade immediately after exit. Confirm strict adherence to {active_rule['name']}: {clean_action}.",
                "[   ] Complete",
            ),
            (
                "Week 2 (Days 8–14)",
                f"Mid-Month Audit: Calculate interim adherence rate for first two weeks. Verify that qualifying trigger events ('{clean_trigger}') were handled strictly according to the rule.",
                "[   ] Complete",
            ),
            (
                "Week 3 (Days 15–21)",
                f"Handling Pressure & Drawdowns: Review execution during choppy sessions or losing streaks. Did you consistently maintain your commitment to {active_rule['name']} without deviation?",
                "[   ] Complete",
            ),
            (
                "Week 4 (Days 22–30)",
                "Month-End Consolidation: Export 30-day tradebook and replay in Nubra Trade Replay to evaluate habit consistency and compare pattern recurrence against your baseline.",
                "[   ] Complete",
            ),
        ]
    for idx, (w_lbl, w_desc, w_st) in enumerate(weeks, start=1):
        row = t_week.rows[idx]
        row.cells[0].paragraphs[0].add_run(w_lbl).bold = True
        row.cells[1].paragraphs[0].add_run(w_desc)
        row.cells[2].paragraphs[0].add_run(w_st)

    # Section 4: Baseline & Success Tracking
    doc.add_page_break()
    doc.add_heading("Section 4 — Measurable Success Tracking", level=1)

    p_form = doc.add_paragraph()
    p_form.add_run("Adherence Rate Formula:\n").bold = True
    p_form.add_run("Adherence Rate (%) = (Trigger Events Where Rule Was Followed ÷ Total Qualifying Trigger Events) × 100\n")
    p_form.add_run("Note: If no qualifying trigger events occur during the month, adherence rate is reported as: 'Not applicable — no qualifying events occurred' (do not divide by zero or assume 100%).").italic = True

    t_track = doc.add_table(rows=6, cols=3)
    t_track.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_track.rows[0].cells[0].paragraphs[0].add_run("Tracking Measure").bold = True
    t_track.rows[0].cells[1].paragraphs[0].add_run("Operational Definition").bold = True
    t_track.rows[0].cells[2].paragraphs[0].add_run("Measurement Method").bold = True

    track_rows = [
        ("Number of Qualifying Trigger Events", f"Total count of trading instances where the trigger condition occurs: {active_rule.get('trigger', '')}", "Recorded in daily tracker as 'Trigger Occurred? = Yes'"),
        ("Number of Events Where Rule Followed", f"Count of qualifying trigger instances where the committed rule action was strictly adhered to: {active_rule.get('action', '')}", "Recorded in daily tracker as 'Rule Followed? = Yes'"),
        ("Number of Violations", "Count of qualifying trigger instances where the committed rule was breached or not followed", "Calculated as: Qualifying Trigger Events − Events Where Rule Followed (logged as 'Rule Followed? = No')"),
        ("Rule Adherence Rate", "Calculated adherence percentage using the formula above: (Events Followed ÷ Total Qualifying Events) × 100", "Calculated at month-end. If 0 qualifying trigger events occurred, reported as: 'Not applicable — no qualifying events occurred'"),
        ("Monitoring Period Review", "Separate review of next-month tradebook; evaluated independently without mixing future results into historical baseline", "Re-evaluate tradebook at month-end. No financial profit or loss reduction is promised or guaranteed."),
    ]
    for idx, (f_area, b_val, t_val) in enumerate(track_rows, start=1):
        row = t_track.rows[idx]
        row.cells[0].paragraphs[0].add_run(f_area).bold = True
        row.cells[1].paragraphs[0].add_run(b_val)
        row.cells[2].paragraphs[0].add_run(t_val)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Complete 30-Day Tracking Log Table (All 30 Days)
    doc.add_heading("Complete 30-Day Adherence Log (Daily Journal Worksheets)", level=2)
    doc.add_paragraph("Fill in each trading session below. If no qualifying trigger event occurred on that day, mark 'N/A'.")

    t_log = doc.add_table(rows=31, cols=4)
    t_log.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Enable header repeating across pages and prevent row splitting across page breaks
    header_tr = t_log.rows[0]._tr.get_or_add_trPr()
    header_tr.append(OxmlElement('w:tblHeader'))
    for row in t_log.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))

    log_headers = ["Day / Date", "Trigger Occurred?", "Rule Followed?", "Trader Notes & Emotional State"]
    for i, h in enumerate(log_headers):
        t_log.rows[0].cells[i].paragraphs[0].add_run(h).bold = True

    for r_idx in range(1, 31):
        row = t_log.rows[r_idx]
        row.cells[0].paragraphs[0].add_run(f"Day {r_idx} / Day {r_idx:02d}: ___/___")
        row.cells[1].paragraphs[0].add_run("[ ] Yes   [ ] No")
        row.cells[2].paragraphs[0].add_run("[ ] Yes   [ ] No   [ ] N/A")
        row.cells[3].paragraphs[0].add_run(" ")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Day-30 Reflection
    doc.add_heading("Day-30 Final Reflection", level=2)
    p_refl = doc.add_paragraph()
    p_refl.add_run("1. Final Adherence Rate: Qualifying events followed: [ _____ ] ÷ Total trigger events: [ _____ ] = [         ]% (or [ ] Not applicable — no qualifying events occurred)\n")
    p_refl.add_run("2. Key Challenge: Under what market conditions or emotional state was it most difficult to follow this rule? ___________________\n")
    p_refl.add_run("3. Next Month Commitment: What habit rule or risk boundary will you carry forward into next month? ___________________\n")

    # Section 5: Limitations
    doc.add_heading("Section 5 — Important Limitations & Disclaimers", level=1)
    p_disc = doc.add_paragraph()
    p_disc.add_run(
        "• Educational Self-Monitoring Aid: Nubra Trade Replay is a post-trade review prototype designed solely to assist "
        "in personal journaling and execution self-awareness. It does not provide financial advice, trading signals, or buy/sell recommendations.\n"
        "• No Guarantee of Financial Outcomes: Past behavioral habits and trade correlations do not predict future market results. "
        "Adhering to habit rules cannot guarantee trading profits or protect against market drawdowns.\n"
        "• Dataset Attribution: All figures and findings are derived directly from the currently selected dataset without external "
        "extrapolation. If this dataset is a synthetic demonstration tradebook, it is intended solely for illustrative habit monitoring."
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
