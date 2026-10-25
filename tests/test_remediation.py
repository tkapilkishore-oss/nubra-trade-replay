"""
Focused regression tests for controlled bug remediation:
- Net-based profit factor (BUG-04)
- Scale-aware and sample-aware expiry comparison (BUG-05)
- Completed trade log breakeven filtering (BUG-06)
- Browser-local persistence helpers (BUG-02 & BUG-03)
"""

import pandas as pd
from src.calculator import calculate_trade_pnl, compute_summary_metrics
from src.detectors import compare_expiry_performance, evaluate_all_detectors
from app import compress_csv_data, decompress_csv_data


def test_profit_factor_gross_winner_turned_net_loser():
    """
    BUG-04: Test that a trade with positive gross P&L but high charges
    (making net P&L negative) is accounted for in net losses, not gross gains.
    """
    df = pd.DataFrame([
        {
            "trade_id": "T_WIN",
            "date": "2026-09-01",
            "entry_time": "09:30:00",
            "exit_time": "09:45:00",
            "symbol": "OPT1",
            "quantity": 10,
            "entry_price": 100.0,
            "exit_price": 120.0,
            "charges": 20.0,
            "is_expiry_day": False,
        },  # Gross = 200, Charges = 20, Net = +180
        {
            "trade_id": "T_GROSS_WIN_NET_LOSS",
            "date": "2026-09-01",
            "entry_time": "10:00:00",
            "exit_time": "10:15:00",
            "symbol": "OPT2",
            "quantity": 10,
            "entry_price": 100.0,
            "exit_price": 105.0,
            "charges": 70.0,
            "is_expiry_day": False,
        },  # Gross = +50, Charges = 70, Net = -20
    ])

    metrics = compute_summary_metrics(df)
    assert metrics.winning_trades == 1
    assert metrics.losing_trades == 1
    # Pure net profit factor = net wins / abs(net losses) = 180 / 20 = 9.0
    assert metrics.profit_factor == 9.0


def test_profit_factor_all_losses_and_no_trades():
    """BUG-04: Edge cases: all losses -> 0.0, zero trades -> None."""
    df_losses = pd.DataFrame([
        {
            "trade_id": "T_LOSS",
            "date": "2026-09-01",
            "entry_time": "09:30:00",
            "exit_time": "09:45:00",
            "symbol": "OPT",
            "quantity": 10,
            "entry_price": 100.0,
            "exit_price": 90.0,
            "charges": 10.0,
            "is_expiry_day": False,
        }  # Net = -110
    ])
    metrics_l = compute_summary_metrics(df_losses)
    assert metrics_l.profit_factor == 0.0

    df_empty = pd.DataFrame(columns=[
        "trade_id", "date", "entry_time", "exit_time", "symbol",
        "quantity", "entry_price", "exit_price", "charges", "is_expiry_day"
    ])
    metrics_e = compute_summary_metrics(df_empty)
    assert metrics_e.profit_factor is None


def test_expiry_comparison_scale_invariance():
    """
    BUG-05: Proportional scaling invariance test.
    Multiplying all financial values by 10x or 100x must preserve relative_delta
    and pattern detection state.
    """
    def make_dataset(multiplier: float):
        records = []
        # 5 expiry trades: -200 * multiplier
        for i in range(5):
            records.append({
                "trade_id": f"EXP_{i}",
                "date": f"2026-09-0{i+1}",
                "entry_time": "09:30:00",
                "exit_time": "09:45:00",
                "symbol": "OPT",
                "quantity": 1,
                "entry_price": 300.0 * multiplier,
                "exit_price": 100.0 * multiplier,
                "charges": 0.0,
                "is_expiry_day": True,
            })
        # 5 non-expiry trades: +100 * multiplier
        for i in range(5):
            records.append({
                "trade_id": f"NON_{i}",
                "date": f"2026-09-1{i}",
                "entry_time": "10:30:00",
                "exit_time": "10:45:00",
                "symbol": "OPT",
                "quantity": 1,
                "entry_price": 100.0 * multiplier,
                "exit_price": 200.0 * multiplier,
                "charges": 0.0,
                "is_expiry_day": False,
            })
        return pd.DataFrame(records)

    res_1x = compare_expiry_performance(make_dataset(1.0))
    res_10x = compare_expiry_performance(make_dataset(10.0))
    res_100x = compare_expiry_performance(make_dataset(100.0))

    # All must trigger pattern detection
    assert res_1x["has_pattern"] is True
    assert res_10x["has_pattern"] is True
    assert res_100x["has_pattern"] is True

    # Relative delta must be strictly identical across 1x, 10x, and 100x
    assert res_1x["relative_delta"] == res_10x["relative_delta"] == res_100x["relative_delta"] == 2.0


def test_expiry_comparison_insufficient_sample():
    """
    BUG-05: If either expiry or non-expiry group has fewer than 3 trades,
    it must return has_sufficient_sample = False and status Insufficient Data.
    """
    records = []
    # Only 2 expiry trades (< 3 minimum)
    for i in range(2):
        records.append({
            "trade_id": f"EXP_{i}",
            "date": "2026-09-01",
            "entry_time": "09:30:00",
            "exit_time": "09:45:00",
            "symbol": "OPT",
            "quantity": 1,
            "entry_price": 100.0,
            "exit_price": 50.0,
            "charges": 0.0,
            "is_expiry_day": True,
        })
    # 5 non-expiry trades
    for i in range(5):
        records.append({
            "trade_id": f"NON_{i}",
            "date": "2026-09-02",
            "entry_time": "10:30:00",
            "exit_time": "10:45:00",
            "symbol": "OPT",
            "quantity": 1,
            "entry_price": 100.0,
            "exit_price": 150.0,
            "charges": 0.0,
            "is_expiry_day": False,
        })
    df = pd.DataFrame(records)
    res = compare_expiry_performance(df)

    assert res["has_both_groups"] is True
    assert res["has_sufficient_sample"] is False
    assert res["has_pattern"] is False
    assert "Insufficient group sample size" in res["status_message"]

    evals = evaluate_all_detectors(df)
    assert evals["expiry_pattern"]["status"] == "Insufficient Data"


def test_expiry_comparison_zero_denominator_safe():
    """BUG-05: Handle trade_scale == 0 gracefully without division by zero."""
    records = [
        {
            "trade_id": f"T_{i}",
            "date": "2026-09-01",
            "entry_time": "09:30:00",
            "exit_time": "09:45:00",
            "symbol": "OPT",
            "quantity": 1,
            "entry_price": 100.0,
            "exit_price": 100.0,
            "charges": 0.0,
            "is_expiry_day": (i % 2 == 0),
        }
        for i in range(8)
    ]
    df = pd.DataFrame(records)
    res = compare_expiry_performance(df)
    assert res["trade_scale"] == 0.0
    assert res["relative_delta"] == 0.0
    assert res["has_pattern"] is False


def test_completed_trade_log_filtering():
    """
    BUG-06: Test tradebook outcome filtering across Winning, Losing,
    Breakeven, and All Trades.
    """
    records = [
        {
            "trade_id": "T_WIN1",
            "date": "2026-09-01",
            "entry_time": "09:30:00",
            "exit_time": "09:45:00",
            "symbol": "OPT",
            "quantity": 10,
            "entry_price": 100.0,
            "exit_price": 110.0,
            "charges": 10.0,
            "is_expiry_day": False,
        },  # Net = +90 (WIN)
        {
            "trade_id": "T_LOSS1",
            "date": "2026-09-01",
            "entry_time": "10:00:00",
            "exit_time": "10:15:00",
            "symbol": "OPT",
            "quantity": 10,
            "entry_price": 100.0,
            "exit_price": 95.0,
            "charges": 10.0,
            "is_expiry_day": False,
        },  # Net = -60 (LOSS)
        {
            "trade_id": "T_BE1",
            "date": "2026-09-01",
            "entry_time": "10:30:00",
            "exit_time": "10:45:00",
            "symbol": "OPT",
            "quantity": 10,
            "entry_price": 100.0,
            "exit_price": 101.0,
            "charges": 10.0,
            "is_expiry_day": False,
        },  # Gross = +10, Charges = 10, Net = 0 (BREAKEVEN)
    ]
    df = calculate_trade_pnl(pd.DataFrame(records))

    # All Trades
    all_trades = df.copy()
    assert len(all_trades) == 3

    # Winning Trades Only
    wins = df[df["net_pnl"] > 0]
    assert len(wins) == 1
    assert wins.iloc[0]["trade_id"] == "T_WIN1"

    # Losing Trades Only
    losses = df[df["net_pnl"] < 0]
    assert len(losses) == 1
    assert losses.iloc[0]["trade_id"] == "T_LOSS1"

    # Breakeven Trades Only
    be = df[df["net_pnl"] == 0]
    assert len(be) == 1
    assert be.iloc[0]["trade_id"] == "T_BE1"


def test_browser_persistence_helpers():
    """
    BUG-02 & BUG-03: Test CSV compression, URL-safe encoding,
    and safe corruption handling.
    """
    original_csv = (
        "trade_id,date,entry_time,exit_time,symbol,quantity,entry_price,exit_price,charges,is_expiry_day\n"
        "T001,2026-09-01,09:30:00,09:45:00,NIFTY_25000_CE,50,120.0,135.0,25.0,False\n"
    )
    encoded = compress_csv_data(original_csv)
    assert isinstance(encoded, str)
    assert len(encoded) > 0

    decoded = decompress_csv_data(encoded)
    assert decoded == original_csv

    # Corrupted / invalid input returns None safely without raising an exception
    assert decompress_csv_data("not_valid_base64_or_zlib!!!") is None
    assert decompress_csv_data("") is None

    # Fresh visit session-state defaults & committed-rule persistence roundtrip
    # Executed in an isolated subprocess to prevent Streamlit's runtime script loader
    # from reloading application modules and causing cross-test class identity divergence.
    import subprocess
    import sys
    from pathlib import Path

    app_path = str(Path(__file__).resolve().parent.parent / "app.py")
    data_dir = str(Path(__file__).resolve().parent.parent / "data")
    test_subproc_code = f"""
import io
from pathlib import Path
import docx
import pandas as pd
import pypdf
from streamlit.testing.v1 import AppTest

from src.action_plan import (
    build_report_context,
    generate_action_plan_docx,
    generate_action_plan_pdf,
    get_rule_by_id,
)
from src.calculator import calculate_trade_pnl, compute_summary_metrics
from src.detectors import generate_ranked_findings

# 1. Fresh visit defaults: Persona B (high_activity)
at_fresh = AppTest.from_file(r"{app_path}")
at_fresh.run()
assert not at_fresh.exception
assert at_fresh.session_state.get("selected_persona_key") == "high_activity"
assert at_fresh.session_state.get("selected_rule_id") == "activity_ceiling"
assert at_fresh.session_state.get("rule_confirmed") is False
assert at_fresh.session_state.get("active_dataset_id") == "sample_high_activity"
sb_fresh = [sb for sb in at_fresh.selectbox if "Select One Rule" in sb.label][0]
assert sb_fresh.value == "activity_ceiling"

# 2. Interactive commitment flow on Persona A (rapid_reentry)
at_a = AppTest.from_file(r"{app_path}")
at_a.query_params["persona"] = "rapid_reentry"
at_a.run()
assert not at_a.exception
assert at_a.session_state.get("selected_persona_key") == "rapid_reentry"
assert at_a.session_state.get("rule_confirmed") is False
assert at_a.session_state.get("selected_rule_id") == "cool_off"

# Interactively select alternative rule 'activity_ceiling' in the selectbox
sb_rule_a = [sb for sb in at_a.selectbox if "Select One Rule" in sb.label][0]
assert sb_rule_a.value == "cool_off"
sb_rule_a.select("activity_ceiling")
at_a.run()
assert not at_a.exception
assert sb_rule_a.value == "activity_ceiling"

# Interactively click the habit commitment button
commit_btn = [b for b in at_a.button if "Commit to Monitoring" in b.label][0]
commit_btn.click()
at_a.run()
assert not at_a.exception

# Assert state & query parameters were updated by the commit button
assert at_a.session_state.get("rule_confirmed") is True
assert at_a.session_state.get("selected_rule_id") == "activity_ceiling"
assert at_a.query_params.get("committed") == "1"
assert at_a.query_params.get("rule") == "activity_ceiling"
assert at_a.query_params.get("dataset_id") == "sample_rapid_reentry"

# 3. Simulate browser reload / restore persisted state
persisted_params = dict(at_a.query_params)
at_reloaded = AppTest.from_file(r"{app_path}")
for k, v in persisted_params.items():
    at_reloaded.query_params[k] = v
at_reloaded.run()
assert not at_reloaded.exception
assert at_reloaded.session_state.get("selected_persona_key") == "rapid_reentry"
assert at_reloaded.session_state.get("rule_confirmed") is True
assert at_reloaded.session_state.get("selected_rule_id") == "activity_ceiling"
sb_reloaded = [sb for sb in at_reloaded.selectbox if "Select One Rule" in sb.label][0]
assert sb_reloaded.value == "activity_ceiling"

# Assert action plan download buttons are rendered on the reloaded page
dl_buttons = at_reloaded.download_button
assert len(dl_buttons) >= 2
assert any("PDF" in b.label for b in dl_buttons)
assert any("Word" in b.label or "Editable" in b.label for b in dl_buttons)

# 4. Verify both PDF and DOCX use the committed alternative rule
df_raw = pd.read_csv(Path(r"{data_dir}") / "rapid_reentry.csv")
df_trades = calculate_trade_pnl(df_raw)
metrics = compute_summary_metrics(df_trades)
findings = generate_ranked_findings(df_trades)
committed_rule = get_rule_by_id(at_reloaded.session_state.selected_rule_id)

ctx = build_report_context(
    dataset_label="Persona A (Rapid Re-entry)",
    is_synthetic=True,
    metrics=metrics,
    findings=findings,
    active_rule=committed_rule,
    df_trades=df_trades,
)
assert ctx.explanation_mode == "user_custom"

pdf_bytes = generate_action_plan_pdf(ctx)
docx_bytes = generate_action_plan_docx(ctx)

pdf_text = " ".join(" ".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(pdf_bytes)).pages).split())
assert "Daily Activity Checkpoint" in pdf_text
assert "However, you have chosen to commit to Daily Activity Checkpoint" in pdf_text
assert "Rapid Re-entry" in pdf_text

doc = docx.Document(io.BytesIO(docx_bytes))
docx_text = " ".join(p.text for p in doc.paragraphs) + " " + " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
assert "Daily Activity Checkpoint" in docx_text
assert "However, you have chosen to commit to Daily Activity Checkpoint" in docx_text
assert "Rapid Re-entry" in docx_text

# 5. Immediate single-interaction persona switching & committed rule isolation
at_switch = AppTest.from_file(r"{app_path}")
at_switch.query_params["persona"] = "rapid_reentry"
at_switch.run()
assert at_switch.session_state.get("selected_persona_key") == "rapid_reentry"
assert at_switch.session_state.get("active_dataset_id") == "sample_rapid_reentry"
full_md_a = " ".join(m.value for m in at_switch.markdown)
assert "Rapid Re-entry Identified (12 Trades)" in full_md_a

# Commit an alternative rule (activity_ceiling) on Persona A
sb_r = [s for s in at_switch.selectbox if "Select One Rule" in s.label][0]
sb_r.select("activity_ceiling")
at_switch.run()
[b for b in at_switch.button if "Commit to Monitoring" in b.label][0].click()
at_switch.run()
assert at_switch.session_state.get("rule_confirmed") is True
assert at_switch.session_state.get("selected_rule_id") == "activity_ceiling"

# Step A: Switch from committed Persona A to Persona B with EXACTLY ONE selectbox interaction
sb_persona = [s for s in at_switch.selectbox if "Select Sample Trader Persona" in s.label][0]
sb_persona.select("high_activity")
at_switch.run()
assert not at_switch.exception

# Assert all Persona B elements updated immediately on that single rerun (metrics, findings, dataset ID, rule)
assert at_switch.session_state.get("selected_persona_key") == "high_activity"
assert at_switch.session_state.get("active_dataset_id") == "sample_high_activity"
assert at_switch.session_state.get("selected_rule_id") == "activity_ceiling"
assert at_switch.session_state.get("rule_confirmed") is False
assert at_switch.query_params.get("persona") == "high_activity"
assert at_switch.query_params.get("committed") == "0"
assert at_switch.query_params.get("rule") is None

full_md_b = " ".join(m.value for m in at_switch.markdown)
assert "High-Activity Losing Days (2 Days)" in full_md_b
assert "-₹2,400.00" in full_md_b
assert "Rapid Re-entry Identified (12 Trades)" not in full_md_b

sb_rule_b = [s for s in at_switch.selectbox if "Select One Rule" in s.label][0]
assert sb_rule_b.value == "activity_ceiling"

# Step B: Switch to additional Persona C (expiry_pattern) with exactly one interaction in the same session
sb_persona = [s for s in at_switch.selectbox if "Select Sample Trader Persona" in s.label][0]
sb_persona.select("expiry_pattern")
at_switch.run()
assert not at_switch.exception
assert at_switch.session_state.get("selected_persona_key") == "expiry_pattern"
assert at_switch.session_state.get("active_dataset_id") == "sample_expiry_pattern"
assert at_switch.session_state.get("selected_rule_id") == "expiry_sizing"
assert at_switch.session_state.get("rule_confirmed") is False
sb_rule_c = [s for s in at_switch.selectbox if "Select One Rule" in s.label][0]
assert sb_rule_c.value == "expiry_sizing"

# Step C: Switch to additional Persona D (clean_trader) with exactly one interaction in the same session
sb_persona = [s for s in at_switch.selectbox if "Select Sample Trader Persona" in s.label][0]
sb_persona.select("clean_trader")
at_switch.run()
assert not at_switch.exception
assert at_switch.session_state.get("selected_persona_key") == "clean_trader"
assert at_switch.session_state.get("active_dataset_id") == "sample_clean_trader"
assert at_switch.session_state.get("selected_rule_id") == "maintain_discipline"
assert at_switch.session_state.get("rule_confirmed") is False
sb_rule_d = [s for s in at_switch.selectbox if "Select One Rule" in s.label][0]
assert sb_rule_d.value == "maintain_discipline"

# Step D: Switch to additional Persona E (tiny_sample) with exactly one interaction in the same session
sb_persona = [s for s in at_switch.selectbox if "Select Sample Trader Persona" in s.label][0]
sb_persona.select("tiny_sample")
at_switch.run()
assert not at_switch.exception
assert at_switch.session_state.get("selected_persona_key") == "tiny_sample"
assert at_switch.session_state.get("active_dataset_id") == "sample_tiny_sample"
assert at_switch.session_state.get("selected_rule_id") == "maintain_discipline"
assert at_switch.session_state.get("rule_confirmed") is False
sb_rule_e = [s for s in at_switch.selectbox if "Select One Rule" in s.label][0]
assert sb_rule_e.value == "maintain_discipline"

# Step E: Verify the switched persona cleanly survives a browser refresh
at_reload_p = AppTest.from_file(r"{app_path}")
for k, v in at_switch.query_params.items():
    at_reload_p.query_params[k] = v
at_reload_p.run()
assert not at_reload_p.exception
assert at_reload_p.session_state.get("selected_persona_key") == "tiny_sample"
assert at_reload_p.session_state.get("rule_confirmed") is False
assert at_reload_p.session_state.get("selected_rule_id") == "maintain_discipline"

# Step F: Stale dataset mismatch URL does not inherit committed state or stale rule
at_mismatch = AppTest.from_file(r"{app_path}")
at_mismatch.query_params["persona"] = "high_activity"
at_mismatch.query_params["committed"] = "1"
at_mismatch.query_params["dataset_id"] = "sample_rapid_reentry"
at_mismatch.query_params["rule"] = "activity_ceiling"
at_mismatch.run()
assert not at_mismatch.exception
assert at_mismatch.session_state.get("selected_persona_key") == "high_activity"
assert at_mismatch.session_state.get("active_dataset_id") == "sample_high_activity"
assert at_mismatch.session_state.get("rule_confirmed") is False
assert at_mismatch.query_params.get("committed") == "0"
assert at_mismatch.query_params.get("rule") is None
assert at_mismatch.query_params.get("dataset_id") == "sample_high_activity"
sb_rule_mismatch = [s for s in at_mismatch.selectbox if "Select One Rule" in s.label][0]
assert sb_rule_mismatch.value == "activity_ceiling"
assert at_mismatch.session_state.get("selected_rule_id") == "activity_ceiling"
"""
    res = subprocess.run([sys.executable, "-c", test_subproc_code], capture_output=True, text=True)
    assert res.returncode == 0, f"AppTest session-state regression failed: {res.stderr}"
