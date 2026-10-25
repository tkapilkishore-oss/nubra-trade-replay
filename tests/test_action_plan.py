"""
Automated unit tests for Nubra Trade Replay Action Plan Document Generation (PDF & DOCX).
Validates document validity, dataset accuracy, disclosures, rule respect, evidence tables,
and adherence log completeness across all personas.
"""

import io
from pathlib import Path
import pandas as pd
import docx

from src.action_plan import (
    ReportContext,
    build_report_context,
    generate_action_plan_docx,
    generate_action_plan_pdf,
)
from src.calculator import calculate_trade_pnl, compute_summary_metrics
from src.detectors import generate_ranked_findings
from src.rules import RULES_CATALOG, get_rule_by_id


DATA_DIR = Path(__file__).resolve().parent.parent / "data"


def load_dataset(filename: str):
    df_raw = pd.read_csv(DATA_DIR / filename)
    df_trades = calculate_trade_pnl(df_raw)
    metrics = compute_summary_metrics(df_trades)
    findings = generate_ranked_findings(df_trades)
    return df_trades, metrics, findings


def test_action_plan_pdf_valid_structure():
    """Verify that PDF generation produces a valid, non-empty PDF document."""
    _, metrics, findings = load_dataset("rapid_reentry.csv")
    rule = get_rule_by_id("cool_off")

    pdf_bytes = generate_action_plan_pdf(
        dataset_label="Persona A (Rapid Re-entry)",
        is_synthetic=True,
        metrics=metrics,
        findings=findings,
        active_rule=rule,
    )

    assert isinstance(pdf_bytes, bytes), "PDF output must be bytes"
    assert len(pdf_bytes) > 20000, f"PDF file size too small: {len(pdf_bytes)} bytes"
    assert pdf_bytes.startswith(b"%PDF-"), "Invalid PDF signature: missing '%PDF-'"
    assert b"%%EOF" in pdf_bytes[-1024:], "PDF missing %%EOF marker"


def test_action_plan_docx_valid_structure():
    """Verify that DOCX generation produces a valid, non-empty Word document."""
    _, metrics, findings = load_dataset("rapid_reentry.csv")
    rule = get_rule_by_id("cool_off")

    docx_bytes = generate_action_plan_docx(
        dataset_label="Persona A (Rapid Re-entry)",
        is_synthetic=True,
        metrics=metrics,
        findings=findings,
        active_rule=rule,
    )

    assert isinstance(docx_bytes, bytes), "DOCX output must be bytes"
    assert len(docx_bytes) > 15000, f"DOCX file size too small: {len(docx_bytes)} bytes"
    assert docx_bytes.startswith(b"PK\x03\x04"), "Invalid DOCX signature (missing PK zip header)"

    doc = docx.Document(io.BytesIO(docx_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    assert len(paragraphs) > 10, "DOCX has too few paragraphs"
    full_text = "\n".join(paragraphs)
    assert "30-Day Trading Behaviour Action Plan" in full_text
    assert "15-Minute Post-Loss Buffer" in full_text


def test_action_plan_dataset_matching_and_metrics():
    """Verify metrics and currency formatting reflect the exact active dataset."""
    _, metrics, findings = load_dataset("rapid_reentry.csv")
    rule = get_rule_by_id("cool_off")

    docx_bytes = generate_action_plan_docx(
        dataset_label="Persona A (Rapid Re-entry)",
        is_synthetic=True,
        metrics=metrics,
        findings=findings,
        active_rule=rule,
    )
    doc = docx.Document(io.BytesIO(docx_bytes))

    table_texts = [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
    combined_tables = " ".join(table_texts)

    # Persona A has 24 trades, net P&L -₹600.00
    assert "24" in combined_tables, "Completed trade count 24 missing from table"
    assert "-₹600.00" in combined_tables, "Net P&L -₹600.00 missing from table"
    assert "15-Minute Post-Loss Buffer" in combined_tables


def test_action_plan_persona_a_content_integrity():
    """Verify Persona A correctly identifies 12 rapid re-entries, evidence, and 15-min rule."""
    _, metrics, findings = load_dataset("rapid_reentry.csv")
    rule = get_rule_by_id("cool_off")

    docx_bytes = generate_action_plan_docx(
        dataset_label="Persona A (Rapid Re-entry)",
        is_synthetic=True,
        metrics=metrics,
        findings=findings,
        active_rule=rule,
    )
    doc = docx.Document(io.BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)

    # Crucial checks per user requirement:
    assert "No recurring behavioral flags were triggered" not in full_text, "False claim that no flags triggered found!"
    assert "Rapid Re-entry Identified (12 Trades)" in full_text, "Rapid re-entry title missing!"
    assert "Preceding Loss Net P&L" in full_text, "Evidence table missing preceding loss column!"
    assert "Flagged Trade Net P&L" in full_text, "Evidence table missing flagged trade column!"
    assert "for 5 minutes" not in full_text, "Contradictory 5 minutes break found!"
    assert " 5 minutes" not in full_text, "Contradictory 5 minutes break found!"
    assert "Wait at least 15 minutes" in full_text, "15-minute wait action missing!"


def test_action_plan_persona_b_and_c_findings_and_evidence():
    """Verify Persona B (High-Activity) and Persona C (Expiry) generate their respective evidence."""
    # Persona B
    _, metrics_b, findings_b = load_dataset("high_activity.csv")
    docx_b = generate_action_plan_docx("Persona B (High-Activity)", True, metrics_b, findings_b, get_rule_by_id("activity_ceiling"))
    text_b = "\n".join(p.text for p in docx.Document(io.BytesIO(docx_b)).paragraphs) + "\n" + "\n".join(c.text for t in docx.Document(io.BytesIO(docx_b)).tables for r in t.rows for c in r.cells)
    assert "High-Activity Losing Days (2 Days)" in text_b
    assert "Daily Trades Executed" in text_b
    assert "12 trades" in text_b

    # Persona C
    _, metrics_c, findings_c = load_dataset("expiry_pattern.csv")
    docx_c = generate_action_plan_docx("Persona C (Expiry Pattern)", True, metrics_c, findings_c, get_rule_by_id("expiry_sizing"))
    text_c = "\n".join(p.text for p in docx.Document(io.BytesIO(docx_c)).paragraphs) + "\n" + "\n".join(c.text for t in docx.Document(io.BytesIO(docx_c)).tables for r in t.rows for c in r.cells)
    assert "Expiry vs Non-Expiry Performance Divergence" in text_c
    assert "Expiry Day Trades" in text_c
    assert "Avg Net P&L / Trade" in text_c


def test_action_plan_persona_d_clean_baseline():
    """Verify Persona D explicitly reports clean baseline with no false patterns."""
    _, metrics_d, findings_d = load_dataset("clean_trader.csv")
    docx_d = generate_action_plan_docx("Persona D (Clean Baseline)", True, metrics_d, findings_d, get_rule_by_id("maintain_discipline"))
    text_d = "\n".join(p.text for p in docx.Document(io.BytesIO(docx_d)).paragraphs)
    assert "Clean Baseline Review: No configured behavioral patterns met detection criteria" in text_d
    assert "Supporting Evidence:" not in text_d, "Clean trader must not show an evidence table for unflagged patterns"


def test_action_plan_synthetic_and_small_sample_notices():
    """Verify synthetic and small-sample warning disclosures appear conditionally."""
    # 1. Persona E (Tiny sample, 8 trades, synthetic)
    _, metrics_e, findings_e = load_dataset("tiny_sample.csv")
    docx_e = generate_action_plan_docx(
        dataset_label="Persona E (Tiny Sample)",
        is_synthetic=True,
        metrics=metrics_e,
        findings=findings_e,
        active_rule=get_rule_by_id("journal_review"),
    )
    doc_e = docx.Document(io.BytesIO(docx_e))
    text_e = "\n".join(p.text for p in doc_e.paragraphs)

    assert "Synthetic Demonstration Data" in text_e, "Synthetic notice missing for Persona E"
    assert "Low-Confidence Notice (Small Sample)" in text_e, "Small sample notice missing for Persona E"
    assert "8 completed trades" in text_e, "Trade count missing from Persona E sample notice"

    # 2. Uploaded Persona F (30 trades, NOT synthetic)
    root_dir = Path(__file__).resolve().parent.parent
    df_f = pd.read_csv(root_dir / "manual_upload_persona_f.csv")
    trades_f = calculate_trade_pnl(df_f)
    metrics_f = compute_summary_metrics(trades_f)
    findings_f = generate_ranked_findings(trades_f)

    docx_f = generate_action_plan_docx(
        dataset_label="User Uploaded Tradebook (manual_upload_persona_f.csv)",
        is_synthetic=False,
        metrics=metrics_f,
        findings=findings_f,
        active_rule=get_rule_by_id("activity_ceiling"),
    )
    doc_f = docx.Document(io.BytesIO(docx_f))
    text_f = "\n".join(p.text for p in doc_f.paragraphs)

    assert "Synthetic Demonstration Data" not in text_f, "Synthetic notice must NOT appear for user upload"
    assert "Low-Confidence Notice (Small Sample)" not in text_f, "Small sample notice must NOT appear for 30 trades"
    assert "User Uploaded Tradebook" in text_f


def test_action_plan_complete_30_day_adherence_log():
    """Verify all 30 days (Day 01 through Day 30) are present in the tracker."""
    _, metrics_a, findings_a = load_dataset("rapid_reentry.csv")
    docx_a = generate_action_plan_docx("Persona A", True, metrics_a, findings_a, get_rule_by_id("cool_off"))
    doc_a = docx.Document(io.BytesIO(docx_a))
    table_text = " ".join(c.text for t in doc_a.tables for r in t.rows for c in r.cells)

    for day_num in range(1, 31):
        day_str = f"Day {day_num:02d}:"
        assert day_str in table_text, f"{day_str} missing from 30-day adherence tracker!"


def test_action_plan_user_selected_alternative_rule():
    """Verify user-selected alternative rule is respected and accurately explained."""
    _, metrics_a, findings_a = load_dataset("rapid_reentry.csv")
    custom_rule = get_rule_by_id("expiry_sizing")

    docx_custom = generate_action_plan_docx(
        dataset_label="Persona A (Rapid Re-entry)",
        is_synthetic=True,
        metrics=metrics_a,
        findings=findings_a,
        active_rule=custom_rule,
    )
    doc_c = docx.Document(io.BytesIO(docx_custom))
    full_text = "\n".join(p.text for p in doc_c.paragraphs) + " " + " ".join(c.text for t in doc_c.tables for r in t.rows for c in r.cells)

    assert "Expiry Day Risk Adjustment" in full_text, "User-selected rule name missing from plan"
    assert "However, you have chosen to commit to Expiry Day Risk Adjustment" in full_text
    assert "Rapid Re-entry Identified (12 Trades)" in full_text


def test_targeted_bugfix_persona_a_comprehensive():
    """Verify all 6 targeted bug fixes specifically for Persona A in generated action plan."""
    _, metrics_a, findings_a = load_dataset("rapid_reentry.csv")
    rule_cool_off = get_rule_by_id("cool_off")

    docx_bytes = generate_action_plan_docx(
        dataset_label="Persona A (Rapid Re-entry)",
        is_synthetic=True,
        metrics=metrics_a,
        findings=findings_a,
        active_rule=rule_cool_off,
    )
    doc = docx.Document(io.BytesIO(docx_bytes))
    full_paragraphs = "\n".join(p.text for p in doc.paragraphs)
    full_tables = "\n".join(" | ".join(c.text.strip().replace("\n", " ") for c in r.cells) for t in doc.tables for r in t.rows)
    all_text = full_paragraphs + "\n" + full_tables

    # Fix 1: Correctly describe detected findings
    assert "No recurring behavioral flags were triggered" not in all_text
    assert "No severe behavioral clusters" not in all_text
    assert "Rapid Re-entry Identified (12 Trades)" in all_text
    assert "12 trades were entered within 15 minutes" in all_text

    # Verified baseline reference values for Persona A
    assert metrics_a.total_trades == 24
    assert metrics_a.active_days == 12
    assert metrics_a.total_net_pnl == -600.0
    assert metrics_a.total_gross_pnl == -600.0
    assert metrics_a.total_charges == 0.0
    assert metrics_a.win_rate == 50.0
    assert metrics_a.winning_trades == 12
    assert metrics_a.losing_trades == 12
    assert metrics_a.avg_net_pnl == -25.0

    # Fix 2: Rule consistency & no contradictory waiting period
    assert "15-Minute Post-Loss Buffer" in all_text
    assert "Wait at least 15 minutes" in all_text
    assert " 5 minutes" not in all_text, "Contradictory 5 minutes break found!"
    assert "for 5 minutes" not in all_text, "Contradictory 5 minutes break found!"
    assert " 5-minute" not in all_text, "Contradictory 5-minute break found!"
    assert "step away" not in all_text.lower()

    # Fix 4: Genuine event-level evidence with total count and distinct P&L attribution
    assert "Total Flagged Events: 12" in all_text
    assert "Preceding Loss Trade" in all_text
    assert "Preceding Loss Net P&L" in all_text
    assert "Flagged Trade Net P&L" in all_text
    assert "R001" in all_text
    assert "R013" in all_text
    assert "09:40:00" in all_text
    assert "09:45:00" in all_text
    assert "5.0 min" in all_text
    assert "-₹100.00" in all_text
    assert "+₹50.00" in all_text
    assert "Preceding Loss Net P&L reflects the loss that immediately preceded" in all_text

    # Fix 5: Measurable success tracking without invented rigid targets
    assert "Adherence Rate Formula" in all_text
    assert "Not applicable — no qualifying events occurred" in all_text
    assert "0 unforced violations" not in all_text
    assert "≥ 90% adherence" not in all_text
    assert "maintain structured pacing" not in all_text
    assert "Number of Qualifying Trigger Events" in all_text
    assert "Number of Events Where Rule Followed" in all_text
    assert "Number of Violations" in all_text
    assert "Rule Adherence Rate" in all_text

    # Fix 6: Currency formatting consistency
    assert "-₹600.00" in all_text
    assert "₹0.00" in all_text
    assert "-₹25.00" in all_text


def test_targeted_bugfix_all_30_tracker_entries_labels():
    """Verify that exactly 30 numbered daily entries (Day 1 through Day 30) are generated."""
    _, metrics_a, findings_a = load_dataset("rapid_reentry.csv")
    rule_cool_off = get_rule_by_id("cool_off")

    # Verify in DOCX
    docx_bytes = generate_action_plan_docx("Persona A", True, metrics_a, findings_a, rule_cool_off)
    doc = docx.Document(io.BytesIO(docx_bytes))
    docx_table_text = " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)

    for day_num in range(1, 31):
        assert f"Day {day_num}" in docx_table_text, f"Day {day_num} label missing from DOCX tracker"
        assert f"Day {day_num:02d}:" in docx_table_text, f"Day {day_num:02d}: missing from DOCX tracker"

    # Verify in PDF
    pdf_bytes = generate_action_plan_pdf("Persona A", True, metrics_a, findings_a, rule_cool_off)
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF-")
    assert b"%%EOF" in pdf_bytes[-1024:]


def test_targeted_bugfix_cross_persona_personalization():
    """Verify that reports for Personas B, C, D, E, F reflect their own detector outputs without leakage."""
    # Persona B
    _, metrics_b, findings_b = load_dataset("high_activity.csv")
    rule_b = get_rule_by_id("activity_ceiling")
    docx_b = generate_action_plan_docx("Persona B (High-Activity)", True, metrics_b, findings_b, rule_b)
    text_b = "\n".join(p.text for p in docx.Document(io.BytesIO(docx_b)).paragraphs) + " " + " ".join(c.text for t in docx.Document(io.BytesIO(docx_b)).tables for r in t.rows for c in r.cells)
    assert "High-Activity Losing Days (2 Days)" in text_b
    assert "Rapid Re-entry Identified" not in text_b, "Persona A finding leaked into Persona B"
    assert "Daily Activity Checkpoint" in text_b

    # Persona C
    _, metrics_c, findings_c = load_dataset("expiry_pattern.csv")
    rule_c = get_rule_by_id("expiry_sizing")
    docx_c = generate_action_plan_docx("Persona C (Expiry)", True, metrics_c, findings_c, rule_c)
    text_c = "\n".join(p.text for p in docx.Document(io.BytesIO(docx_c)).paragraphs) + " " + " ".join(c.text for t in docx.Document(io.BytesIO(docx_c)).tables for r in t.rows for c in r.cells)
    assert "Expiry vs Non-Expiry Performance Divergence" in text_c
    assert "Expiry Day Risk Adjustment" in text_c
    assert "Rapid Re-entry" not in text_c

    # Persona D (Clean)
    _, metrics_d, findings_d = load_dataset("clean_trader.csv")
    rule_d = get_rule_by_id("maintain_discipline")
    docx_d = generate_action_plan_docx("Persona D (Clean)", True, metrics_d, findings_d, rule_d)
    text_d = "\n".join(p.text for p in docx.Document(io.BytesIO(docx_d)).paragraphs) + " " + " ".join(c.text for t in docx.Document(io.BytesIO(docx_d)).tables for r in t.rows for c in r.cells)
    assert "Clean Baseline Review: No configured behavioral patterns met detection criteria" in text_d
    assert "Supporting Evidence:" not in text_d

    # Persona E (Small sample)
    _, metrics_e, findings_e = load_dataset("tiny_sample.csv")
    rule_e = get_rule_by_id("journal_review")
    docx_e = generate_action_plan_docx("Persona E (Tiny)", True, metrics_e, findings_e, rule_e)
    text_e = "\n".join(p.text for p in docx.Document(io.BytesIO(docx_e)).paragraphs)
    assert "Low-Confidence Notice (Small Sample)" in text_e
    assert "8 completed trades" in text_e

    # Persona F (User upload)
    root_dir = Path(__file__).resolve().parent.parent
    df_f = pd.read_csv(root_dir / "manual_upload_persona_f.csv")
    trades_f = calculate_trade_pnl(df_f)
    metrics_f = compute_summary_metrics(trades_f)
    findings_f = generate_ranked_findings(trades_f)
    docx_f = generate_action_plan_docx("User Uploaded Tradebook", False, metrics_f, findings_f, get_rule_by_id("activity_ceiling"))
    text_f = "\n".join(p.text for p in docx.Document(io.BytesIO(docx_f)).paragraphs)
    assert "User Uploaded Tradebook" in text_f
    assert "Synthetic Demonstration Data" not in text_f


def test_targeted_corrections_persona_b_wording_and_trigger():
    """Verify Persona B report has activity-specific wording, consistent 4th-trade trigger, and clean currency in PDF and DOCX."""
    import pypdf
    _, metrics_b, findings_b = load_dataset("high_activity.csv")
    rule_b = get_rule_by_id("activity_ceiling")

    consistent_trigger = "When the fourth completed trade of the day is reached, pause and reassess before placing another trade."
    activity_narrative = "pausing when the daily trade-count threshold is reached, reassessing market conditions and strategy fit, and avoiding unnecessary additional trades"

    # 1. Test DOCX
    docx_bytes = generate_action_plan_docx("Persona B (High-Activity)", True, metrics_b, findings_b, rule_b)
    doc_b = docx.Document(io.BytesIO(docx_bytes))
    p_text_b = "\n".join(p.text for p in doc_b.paragraphs)
    t_text_b = " ".join(c.text for t in doc_b.tables for r in t.rows for c in r.cells)
    all_text_docx_b = " ".join((p_text_b + " " + t_text_b).split())

    # Trigger consistency across rule card, checklist, milestones, tracking, narrative
    assert consistent_trigger in all_text_docx_b, "Consistent 4th-trade trigger missing in DOCX"
    assert activity_narrative in all_text_docx_b, "Activity-specific narrative missing in DOCX"
    assert "impulse re-entry" not in all_text_docx_b, "Rapid re-entry language leaked into Persona B DOCX"
    assert "independent setup" not in all_text_docx_b, "Independent setup wording leaked into Persona B DOCX"
    assert "-₹2,400.00" in all_text_docx_b, "Formatted -₹2,400.00 missing in Persona B DOCX"
    assert "-₹1,200.00" in all_text_docx_b, "Formatted -₹1,200.00 missing in Persona B DOCX"
    assert "₹-" not in all_text_docx_b, "Malformed ₹- amount found in Persona B DOCX"
    assert ".." not in (p_text_b + " " + t_text_b), "Duplicate punctuation (..) found in Persona B DOCX"

    # Verify complete sentence for high-activity finding summary
    complete_ha_sentence = "The review identified 2 days with unusually high trading activity"
    assert complete_ha_sentence in all_text_docx_b, "Complete high-activity finding sentence missing in DOCX"
    assert "P&L;" not in all_text_docx_b, "Erroneous semicolon in P&L found in Persona B DOCX"
    assert "Net Realized P&L" in all_text_docx_b
    assert "Gross Realized P&L" in all_text_docx_b
    assert "Avg Net P&L / Trade" in all_text_docx_b
    assert "Daily Net Realized P&L" in all_text_docx_b

    # Verify detector threshold and historical findings remain unchanged
    assert "High-Activity Losing Days (2 Days)" in all_text_docx_b
    assert "4.0 trades/day" in all_text_docx_b
    assert "12 trades" in all_text_docx_b

    # 2. Test PDF
    pdf_bytes = generate_action_plan_pdf("Persona B (High-Activity)", True, metrics_b, findings_b, rule_b)
    r_pdf_b = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    raw_pdf_text_b = "\n".join(p.extract_text() for p in r_pdf_b.pages)
    all_text_pdf_b = " ".join(raw_pdf_text_b.split())

    assert consistent_trigger in all_text_pdf_b, "Consistent 4th-trade trigger missing in PDF"
    assert activity_narrative in all_text_pdf_b, "Activity-specific narrative missing in PDF"
    assert "impulse re-entry" not in all_text_pdf_b, "Rapid re-entry language leaked into Persona B PDF"
    assert "independent setup" not in all_text_pdf_b, "Independent setup wording leaked into Persona B PDF"
    assert "-₹2,400.00" in all_text_pdf_b, "Formatted -₹2,400.00 missing in Persona B PDF"
    assert "-₹1,200.00" in all_text_pdf_b, "Formatted -₹1,200.00 missing in Persona B PDF"
    assert "₹-" not in all_text_pdf_b, "Malformed ₹- amount found in Persona B PDF"
    assert ".." not in raw_pdf_text_b, "Duplicate punctuation (..) found in Persona B PDF"
    assert complete_ha_sentence in all_text_pdf_b, "Complete high-activity finding sentence missing in PDF"
    assert "P&L;" not in raw_pdf_text_b, "Erroneous semicolon in P&L found in Persona B PDF"
    assert "Net Realized P&L" in all_text_pdf_b
    assert "Gross Realized P&L" in all_text_pdf_b
    assert "Avg Net P&L / Trade" in all_text_pdf_b
    assert "Daily Net Realized P&L" in all_text_pdf_b
    assert "High-Activity Losing Days (2 Days)" in all_text_pdf_b


def test_targeted_corrections_persona_a_copy_and_formatting():
    """Verify Persona A preserves 15m rule & evidence, uses non-causal opportunity phrasing, and has clean currency/punctuation."""
    import pypdf
    _, metrics_a, findings_a = load_dataset("rapid_reentry.csv")
    rule_a = get_rule_by_id("cool_off")

    opportunity_phrasing = "providing an opportunity to check that the next trade is based on an independent setup rather than an impulse re-entry"
    deduped_explanation = "by establishing a structured pause after a losing exit, providing an opportunity to check that the next trade is based on an independent setup rather than an impulse re-entry"
    duplicated_providing = "by providing a structured pause after a losing exit, providing"
    attribution_note = "Preceding Loss Net P&L reflects the loss that immediately preceded the rapid entry"

    # 1. Test DOCX
    docx_bytes = generate_action_plan_docx("Persona A (Rapid Re-entry)", True, metrics_a, findings_a, rule_a)
    doc_a = docx.Document(io.BytesIO(docx_bytes))
    p_text_a = "\n".join(p.text for p in doc_a.paragraphs)
    t_text_a = " ".join(c.text for t in doc_a.tables for r in t.rows for c in r.cells)
    all_text_docx_a = " ".join((p_text_a + " " + t_text_a).split())

    assert "15-Minute Post-Loss Buffer" in all_text_docx_a
    assert "Wait at least 15 minutes" in all_text_docx_a
    assert "Rapid Re-entry Identified (12 Trades)" in all_text_docx_a
    assert "Total Flagged Events: 12" in all_text_docx_a
    assert "Preceding Loss Net P&L" in all_text_docx_a
    assert "Flagged Trade Net P&L" in all_text_docx_a
    assert opportunity_phrasing in all_text_docx_a
    assert deduped_explanation in all_text_docx_a, "Deduplicated Persona A explanation missing in DOCX"
    assert duplicated_providing not in all_text_docx_a, "Duplicated 'providing' found in Persona A DOCX"
    assert "ensures every subsequent trade" not in all_text_docx_a
    assert "ensure every subsequent trade" not in all_text_docx_a
    assert attribution_note in all_text_docx_a
    assert "-₹600.00" in all_text_docx_a
    assert "₹-" not in all_text_docx_a
    assert ".." not in (p_text_a + " " + t_text_a)
    assert "P&L;" not in all_text_docx_a, "Erroneous semicolon in P&L found in Persona A DOCX"
    assert "Net Realized P&L" in all_text_docx_a
    assert "Gross Realized P&L" in all_text_docx_a
    assert "Avg Net P&L / Trade" in all_text_docx_a

    # 2. Test PDF
    pdf_bytes = generate_action_plan_pdf("Persona A (Rapid Re-entry)", True, metrics_a, findings_a, rule_a)
    r_pdf_a = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    raw_pdf_text_a = "\n".join(p.extract_text() for p in r_pdf_a.pages)
    all_text_pdf_a = " ".join(raw_pdf_text_a.split())

    assert "15-Minute Post-Loss Buffer" in all_text_pdf_a
    assert "Wait at least 15 minutes" in all_text_pdf_a
    assert "Rapid Re-entry Identified (12 Trades)" in all_text_pdf_a
    assert "Total Flagged Events: 12" in all_text_pdf_a
    assert "Preceding Loss Net P&L" in all_text_pdf_a
    assert "Flagged Trade Net P&L" in all_text_pdf_a
    assert opportunity_phrasing in all_text_pdf_a
    assert deduped_explanation in all_text_pdf_a, "Deduplicated Persona A explanation missing in PDF"
    assert duplicated_providing not in all_text_pdf_a, "Duplicated 'providing' found in Persona A PDF"
    assert "ensures every subsequent trade" not in all_text_pdf_a
    assert "ensure every subsequent trade" not in all_text_pdf_a
    assert "-₹600.00" in all_text_pdf_a
    assert "₹-" not in all_text_pdf_a
    assert ".." not in raw_pdf_text_a
    assert "P&L;" not in raw_pdf_text_a, "Erroneous semicolon in P&L found in Persona A PDF"
    assert "Net Realized P&L" in all_text_pdf_a
    assert "Gross Realized P&L" in all_text_pdf_a
    assert "Avg Net P&L / Trade" in all_text_pdf_a


def test_unified_pipeline_seeded_personas_parity():
    """Verify that all seeded personas (A, B, C, D, E) generate consistent 4-page PDFs and full DOCXs via ReportContext."""
    import pypdf
    persona_configs = [
        ("Persona A", "rapid_reentry.csv", "cool_off", "15-Minute Post-Loss Buffer"),
        ("Persona B", "high_activity.csv", "activity_ceiling", "Daily Activity Checkpoint"),
        ("Persona C", "expiry_pattern.csv", "expiry_sizing", "Expiry Day Risk Adjustment"),
        ("Persona D", "clean_trader.csv", "maintain_discipline", "Preserve Current Execution Standards"),
        ("Persona E", "tiny_sample.csv", "journal_review", "Weekly Execution Journal"),
    ]

    canonical_sections = [
        "NUBRA TRADE REPLAY",
        "30-Day Trading Behaviour Action Plan",
        "Selected Habit Rule Commitment",
        "Baseline Portfolio Review Metrics",
        "Section 2 — Why This Rule Was Selected",
        "Section 3 — The 30-Day Action Checklist",
        "Weekly Review Milestones",
        "Section 4 — Measurable Success Tracking",
        "Complete 30-Day Adherence Log",
        "Day-30 Final Reflection",
        "Section 5 — Important Limitations & Disclaimers",
    ]

    for name, csv_file, rule_id, rule_name in persona_configs:
        df_raw, metrics, findings = load_dataset(csv_file)
        rule = get_rule_by_id(rule_id)
        from src.action_plan import format_currency_text
        exp_pnl = format_currency_text(metrics.total_net_pnl, True)

        ctx = build_report_context(
            dataset_label=f"{name} Benchmark",
            is_synthetic=True,
            metrics=metrics,
            findings=findings,
            active_rule=rule,
            df_trades=df_raw,
        )
        assert isinstance(ctx, ReportContext)

        # PDF
        pdf_bytes = generate_action_plan_pdf(ctx)
        r_pdf = pypdf.PdfReader(io.BytesIO(pdf_bytes))
        assert len(r_pdf.pages) == 4, f"{name} PDF must be exactly 4 pages, got {len(r_pdf.pages)}"
        raw_pdf = "\n".join(p.extract_text() for p in r_pdf.pages)
        norm_pdf = " ".join(raw_pdf.split())

        for sec in canonical_sections:
            assert sec in norm_pdf, f"Missing section '{sec}' in {name} PDF"

        assert rule_name in norm_pdf
        assert exp_pnl in norm_pdf
        assert "P&L;" not in raw_pdf
        assert "Synthetic Demonstration Data" in norm_pdf
        assert "Day 1" in norm_pdf and "Day 30" in norm_pdf

        # DOCX
        docx_bytes = generate_action_plan_docx(ctx)
        doc = docx.Document(io.BytesIO(docx_bytes))
        raw_docx = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        norm_docx = " ".join(raw_docx.split())

        for sec in canonical_sections:
            assert sec in norm_docx, f"Missing section '{sec}' in {name} DOCX"

        assert rule_name in norm_docx
        assert exp_pnl in norm_docx
        assert "P&L;" not in raw_docx
        assert "Day 30" in norm_docx


def test_unified_pipeline_uploaded_csv():
    """Verify that a custom uploaded CSV uses the exact same canonical pipeline and template as seeded personas."""
    import pypdf
    from src.validator import parse_and_validate_csv

    root_dir = Path(__file__).resolve().parent.parent
    csv_path = root_dir / "manual_upload_persona_f.csv"
    val_res = parse_and_validate_csv(csv_path)
    assert val_res.is_valid

    df_trades = calculate_trade_pnl(val_res.df)
    metrics = compute_summary_metrics(df_trades)
    findings = generate_ranked_findings(df_trades)
    active_rule = get_rule_by_id("activity_ceiling")

    # Build report context
    ctx = build_report_context(
        dataset_label="Custom Uploaded Tradebook (manual_upload_persona_f.csv)",
        is_synthetic=False,
        metrics=metrics,
        findings=findings,
        active_rule=active_rule,
        df_trades=df_trades,
    )

    canonical_sections = [
        "NUBRA TRADE REPLAY",
        "30-Day Trading Behaviour Action Plan",
        "Selected Habit Rule Commitment",
        "Baseline Portfolio Review Metrics",
        "Section 2 — Why This Rule Was Selected",
        "Section 3 — The 30-Day Action Checklist",
        "Weekly Review Milestones",
        "Section 4 — Measurable Success Tracking",
        "Complete 30-Day Adherence Log",
        "Day-30 Final Reflection",
        "Section 5 — Important Limitations & Disclaimers",
    ]

    # Generate PDF
    pdf_bytes = generate_action_plan_pdf(ctx)
    r_pdf = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    assert len(r_pdf.pages) == 4, f"Uploaded CSV PDF must be 4 pages, got {len(r_pdf.pages)}"
    raw_pdf = "\n".join(p.extract_text() for p in r_pdf.pages)
    norm_pdf = " ".join(raw_pdf.split())

    for sec in canonical_sections:
        assert sec in norm_pdf, f"Missing section '{sec}' in uploaded CSV PDF"

    # Verify custom upload content
    assert "User Uploaded Tradebook" in norm_pdf
    assert "Synthetic Demonstration Data" not in norm_pdf
    assert "30" in norm_pdf  # 30 completed trades
    assert "-₹9,275.00" in norm_pdf
    assert "High-Activity Losing Days (2 Days)" in norm_pdf
    assert "P&L;" not in raw_pdf
    assert "Day 1" in norm_pdf and "Day 30" in norm_pdf
    # Zero leakage from Persona A
    assert "15-Minute Post-Loss Buffer" not in norm_pdf
    assert "impulse re-entry" not in norm_pdf

    # Generate DOCX
    docx_bytes = generate_action_plan_docx(ctx)
    doc = docx.Document(io.BytesIO(docx_bytes))
    raw_docx = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    norm_docx = " ".join(raw_docx.split())

    for sec in canonical_sections:
        assert sec in norm_docx, f"Missing section '{sec}' in uploaded CSV DOCX"

    assert "User Uploaded Tradebook" in norm_docx
    assert "Synthetic Demonstration Data" not in norm_docx
    assert "-₹9,275.00" in norm_docx
    assert "Day 30" in norm_docx
    assert "P&L;" not in raw_docx


def test_unified_pipeline_dataset_switching_and_isolation():
    """Verify that switching datasets in sequence (Persona A -> Persona B -> Uploaded CSV -> Persona A) maintains total state isolation."""
    import pypdf
    from src.validator import parse_and_validate_csv

    root_dir = Path(__file__).resolve().parent.parent

    # Step 1: Persona A
    df_a, metrics_a, findings_a = load_dataset("rapid_reentry.csv")
    ctx_a = build_report_context("Persona A", True, metrics_a, findings_a, get_rule_by_id("cool_off"), df_a)
    pdf_a = generate_action_plan_pdf(ctx_a)
    text_a = " ".join("\n".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(pdf_a)).pages).split())
    assert "Rapid Re-entry Identified (12 Trades)" in text_a
    assert "High-Activity Losing Days" not in text_a

    # Step 2: Switch to Persona B
    df_b, metrics_b, findings_b = load_dataset("high_activity.csv")
    ctx_b = build_report_context("Persona B", True, metrics_b, findings_b, get_rule_by_id("activity_ceiling"), df_b)
    pdf_b = generate_action_plan_pdf(ctx_b)
    text_b = " ".join("\n".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(pdf_b)).pages).split())
    assert "High-Activity Losing Days (2 Days)" in text_b
    assert "Rapid Re-entry Identified" not in text_b
    assert "-₹1,600.00" in text_b

    # Step 3: Switch to Uploaded CSV
    val_f = parse_and_validate_csv(root_dir / "manual_upload_persona_f.csv")
    trades_f = calculate_trade_pnl(val_f.df)
    ctx_f = build_report_context("User Uploaded", False, compute_summary_metrics(trades_f), generate_ranked_findings(trades_f), get_rule_by_id("activity_ceiling"), trades_f)
    pdf_f = generate_action_plan_pdf(ctx_f)
    text_f = " ".join("\n".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(pdf_f)).pages).split())
    assert "-₹9,275.00" in text_f
    assert "User Uploaded" in text_f
    assert "Synthetic Demonstration Data" not in text_f

    # Step 4: Switch back to Persona A
    ctx_a2 = build_report_context("Persona A", True, metrics_a, findings_a, get_rule_by_id("cool_off"), df_a)
    pdf_a2 = generate_action_plan_pdf(ctx_a2)
    text_a2 = " ".join("\n".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(pdf_a2)).pages).split())
    assert "Rapid Re-entry Identified (12 Trades)" in text_a2
    assert "-₹9,275.00" not in text_a2
    assert "High-Activity Losing Days" not in text_a2


def test_unified_pipeline_edge_cases():
    """Verify edge cases: clean baseline, tiny sample warning, user custom rule, and adherence formula zero-trigger handling."""
    import pypdf

    # Case 1: Clean baseline trader (Persona D)
    df_d, metrics_d, findings_d = load_dataset("clean_trader.csv")
    ctx_d = build_report_context("Clean Trader", True, metrics_d, findings_d, get_rule_by_id("maintain_discipline"), df_d)
    assert ctx_d.explanation_mode == "clean"
    pdf_d = generate_action_plan_pdf(ctx_d)
    text_d = " ".join("\n".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(pdf_d)).pages).split())
    assert "Clean Baseline Review: No configured behavioral patterns met detection criteria" in text_d
    assert "Supporting Evidence:" not in text_d

    # Case 2: Tiny sample (<10 trades) (Persona E)
    df_e, metrics_e, findings_e = load_dataset("tiny_sample.csv")
    ctx_e = build_report_context("Tiny Sample", True, metrics_e, findings_e, get_rule_by_id("journal_review"), df_e)
    assert ctx_e.is_small_sample is True
    pdf_e = generate_action_plan_pdf(ctx_e)
    text_e = " ".join("\n".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(pdf_e)).pages).split())
    assert "Low-Confidence Notice (Small Sample)" in text_e
    assert "8 completed trades" in text_e

    # Case 3: User custom rule commitment (e.g. committing to expiry_sizing on Persona A)
    df_a, metrics_a, findings_a = load_dataset("rapid_reentry.csv")
    custom_rule = get_rule_by_id("expiry_sizing")
    ctx_custom = build_report_context("Custom Choice", True, metrics_a, findings_a, custom_rule, df_a)
    assert ctx_custom.explanation_mode == "user_custom"
    pdf_custom = generate_action_plan_pdf(ctx_custom)
    text_custom = " ".join("\n".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(pdf_custom)).pages).split())
    assert "The primary behavioral pattern detected in your review was Rapid Re-entry Identified" in text_custom
    assert "However, you have chosen to commit to Expiry Day Risk Adjustment" in text_custom
    assert "Supporting Evidence: Rapid Re-entry Identified" in text_custom

    # Case 4: Zero-trigger adherence formula in definition
    assert "Not applicable — no qualifying events occurred" in text_d
    assert "Not applicable — no qualifying events occurred" in text_custom


